"""Datasety Brány 1 a 2: stiahnutie, import do DuckDB, profil, pokrytie organizácií a bezpečné dotazy.

- Súbor sa sťahuje po častiach s limitom veľkosti (gate1.max_dataset_mb) a celkovým limitom
  (gate1.max_total_gb); blokované domény (sociálne siete) sa nesťahujú ani po presmerovaní, robots.txt platí.
- Tabuľky sú v jednej DuckDB databáze na hypotézu (data/datasets/H<id>/gate1.duckdb, mimo gitu); pôvodný
  súbor sa po importe zmaže (ostáva URL, čas a SHA-256 – vieme, čo presne sa stiahlo).
- Dataset s osobnými údajmi (stĺpce meno, rodné číslo, e-mail… alebo e-maily a telefóny v hodnotách) sa
  neuloží: najprv posúdenie podľa GDPR (R-PRIV). Ostane iba profil bez ukážok.
- Široké tabuľky (ŠÚ SR DATAcube, Eurostat: roky, štvrťroky alebo mesiace ako stĺpce) dostanú aj dlhý tvar
  `dsN_long` (identifikačné stĺpce, obdobie DATE, obdobie_stlpec, hodnota DOUBLE) – časový rad pre Bránu 2.
- Dotazy od modelu (reprodukcia čísla) bežia iba na čítanie, bez prístupu k súborom a sieti (DuckDB
  enable_external_access = false), iba jeden SELECT a s časovým limitom. Čísla počíta DuckDB/Python.
"""

from __future__ import annotations

import hashlib
from collections import Counter
from io import BytesIO
import logging
import re
import threading
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import duckdb
import phonenumbers

from .config import Config
from .fetch import domain_blocked

log = logging.getLogger(__name__)

FILE_EXT = ("csv", "tsv", "txt", "xlsx", "json", "parquet", "zip", "xml")
FILE_LINK_RE = re.compile(r"\.(csv|tsv|xlsx|json|parquet|zip|xml)(\?|$)", re.I)
TABLE_FORMATS = ("csv", "tsv", "txt", "xlsx", "json", "parquet", "xml")
PERSONAL_COLUMN_RE = re.compile(
    r"(^|_|\b)(meno|priezvisko|rodn[eé]|rod_?c|r\.?\s?č\.?|e-?mail|telef|mobil|adres|ulica|bydlisk|narod|"
    r"name|surname|birth|phone|address)", re.I)
EMAIL_VALUE_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_VALUE_RE = re.compile(r"^\+?\d[\d /]{8,}$")
ICO_COLUMN_RE = re.compile(r"(^|_|\b)(i[čc]o|ic)(\b|_|$)", re.I)
TIME_COLUMN_RE = re.compile(r"(rok|year|d[áa]tum|date|obdob|mesiac|month|kvart|quarter|time|čas)", re.I)
TABLE_RE = re.compile(r"ds\d+(_long)?")  # tabuľka datasetu a jej dlhý tvar
_Y = r"(?P<y>19[5-9]\d|20\d\d)"
_UNIT = r"(?:\s*[\[(].*[\])])?"  # „2020 [EUR]“, „2020 (počet)“
PERIOD_NAME_RES = [  # obdobie v názve stĺpca širokej tabuľky
    ("quarter", re.compile(rf"^{_Y}\s*[-_/ ]?\s*Q(?P<q>[1-4]){_UNIT}$", re.I)),  # 2021-Q1, 2021Q1, 2021 Q1
    ("quarter", re.compile(rf"^Q(?P<q>[1-4])\s*[-_/ ]?\s*{_Y}{_UNIT}$", re.I)),  # Q1 2021
    ("quarter", re.compile(rf"^(?P<q>[1-4])\.?\s*(?:Q|štvrťrok|stvrtrok|kvartál|kvartal)\s*{_Y}{_UNIT}$", re.I)),
    ("month", re.compile(rf"^{_Y}\s*(?:-|/|M|_)\s*(?P<m>0?[1-9]|1[0-2]){_UNIT}$", re.I)),  # 2022-01, 2022M01
    ("month", re.compile(rf"^(?P<m>0?[1-9]|1[0-2])\s*[/.]\s*{_Y}{_UNIT}$")),  # 01/2022, 1.2022
    ("year", re.compile(rf"^(?:rok\s*|year\s*)?{_Y}{_UNIT}$", re.I)),  # 2020, rok 2020, 2020 [EUR]
]
FORBIDDEN_SQL = re.compile(
    r"\b(read_\w+|copy|attach|detach|install|load|pragma|export|import|create|insert|update|delete|drop|alter|"
    r"call|set|reset|checkpoint|vacuum|glob|httpfs|system)\b|https?://|;.+", re.I | re.S)


def _is_phone(value: str) -> bool:
    """Telefón v hodnote: začína + alebo 0 a phonenumbers ho uzná (DIČ, IČO a sumy nie sú telefóny)."""
    v = value.strip()
    if not PHONE_VALUE_RE.match(v) or not v.startswith(("+", "0")):
        return False
    try:
        return phonenumbers.is_valid_number(phonenumbers.parse(v, "SK"))
    except phonenumbers.NumberParseException:
        return False


class DatasetError(RuntimeError):
    pass


@dataclass
class Download:
    path: Path
    sha256: str
    size: int
    final_url: str
    content_type: str
    note: str | None = None  # napr. počet strán API alebo skrátenie na limit
    window: tuple[str, str] | None = None  # obdobie, ktoré stiahnuté dáta pokrývajú (napr. dni dávok CRZ)


SHARED_MARKER = "zdielane_z.txt"  # podhypotéza používa DuckDB rodiča (rozdelenie hypotézy)



def connect_ro(path: Path | str, config: dict | None = None, attempts: int = 12, wait_s: float = 5.0):
    """DuckDB iba na čítanie; keď ju práve zapisuje iný proces (doplňovanie CRZ, krok datasetov), chvíľu počká."""
    import time

    for i in range(attempts):
        try:
            return duckdb.connect(str(path), read_only=True, config=config or {})
        except duckdb.IOException:
            if i == attempts - 1:
                raise
            time.sleep(wait_s)

def hyp_dir(cfg: Config, hid: int) -> Path:
    base = cfg.path(cfg.gate1.datasets_dir)
    own = base / f"H{hid}"
    marker = own / SHARED_MARKER
    if marker.exists():
        try:
            return base / f"H{int(marker.read_text(encoding='utf-8').strip())}"
        except ValueError:
            pass
    return own


def share_datasets(cfg: Config, child: int, parent: int) -> None:
    """Podhypotéza číta a zapisuje do DuckDB rodiča (tabuľky majú globálne jedinečné čísla dsN)."""
    root = hyp_dir(cfg, parent)
    own = cfg.path(cfg.gate1.datasets_dir) / f"H{child}"
    own.mkdir(parents=True, exist_ok=True)
    (own / SHARED_MARKER).write_text(root.name[1:], encoding="utf-8")


def duck_path(cfg: Config, hid: int) -> Path:
    return hyp_dir(cfg, hid) / "gate1.duckdb"


def file_format(url: str, content_type: str = "") -> str | None:
    path = urlparse(url).path.lower()
    for ext in FILE_EXT:
        if path.endswith("." + ext):
            return "csv" if ext in ("tsv", "txt") else ext
    ct = content_type.lower()
    if "csv" in ct:
        return "csv"
    if "spreadsheetml" in ct:
        return "xlsx"
    if "json" in ct:
        return "json"
    if "parquet" in ct:
        return "parquet"
    if "zip" in ct:
        return "zip"
    if "xml" in ct and "html" not in ct:
        return "xml"
    return None


PARSER_VERSION = "2026-10-15"  # verzia importu (adaptéry, hlavičky, kontroly) – zapíše sa do profilu datasetu
# Stĺpce formátovania buniek a grafu (SpreadsheetML, RDF), ktoré nie sú významové dáta (audit: H46/ds36–38, ds31).
FORMAT_COLUMNS = {"styleid", "mergeacross", "mergedown", "formula", "index", "type", "data", "cell", "row", "nodeid",
                  "about", "resource", "hasvalue", "datatype", "lang", "href", "hreftooltip", "arrayrange", "ticked"}
# jednotka v názve stĺpca; podčiarkovník je oddeľovač („suma_eur“), „hodnota“ nie je „hod“
UNIT_RE = re.compile(r"(?<![^\W_])(€|eur|czk|ks|%|počet|tis|mil|kg|m2|dni|hod|hodín|hodiny)(?![^\W_])", re.I)


def sniff_format(path: Path, fmt: str) -> str:
    """Formát podľa skutočného obsahu, nie podľa prípony: ZIP/XLSX, Parquet, SpreadsheetML, RDF, JSON-LD, XML, JSON.
    Stránka HTML (prihlásenie, chyba, rozcestník) sa odmietne – nie je to tabuľka (audit: ds52 s hlavičkou DOCTYPE)."""
    from .fetch import looks_like_html

    with open(path, "rb") as fh:
        head = fh.read(4096)
    if head.startswith(b"PK\x03\x04"):
        return "xlsx" if fmt == "xlsx" else "zip"
    if head.startswith(b"PAR1"):
        return "parquet"
    if looks_like_html(head):
        raise DatasetError("namiesto dát stránka HTML (prihlásenie, chyba alebo rozcestník) – nie je to tabuľka")
    text = head.lstrip(b"\xef\xbb\xbf").lstrip()
    low = head.lower()
    if text.startswith(b"<"):
        if b"urn:schemas-microsoft-com:office:spreadsheet" in low:
            return "spreadsheetml"
        if b"22-rdf-syntax-ns" in low or b"<rdf:rdf" in low:
            return "rdf"
        return "xml"
    if text.startswith((b"{", b"[")) and fmt in ("json", "bin", "txt", "csv"):
        return "jsonld" if (b'"@context"' in head or b'"@graph"' in head) else "json"
    return fmt


# ---------------------------------------------------------------- tabuľka vo výkaze (XLSX, SpreadsheetML, CSV)
def _is_number(v: Any) -> bool:
    if isinstance(v, (int, float)):
        return True
    t = str(v).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    return bool(re.fullmatch(r"[-+]?\d+(\.\d+)?%?", t))


def _is_texty(v: Any) -> bool:
    """Bunka vyzerá ako názov stĺpca: text, nie číslo ani dátum."""
    if v is None or str(v).strip() == "":
        return False
    if _is_number(v) or isinstance(v, (date,)):
        return False
    return not re.fullmatch(r"\d{1,4}[./-]\d{1,2}([./-]\d{1,4})?", str(v).strip())


def _is_year(v: Any) -> bool:
    """Rok v druhom riadku hlavičky („výdavky“ nad „2020 | 2021“) – číslo, no názov stĺpca."""
    return bool(re.fullmatch(r"(19|20)\d\d(\.0)?", str(v).strip()))


def detect_table(rows: list[list[Any]], scan: int = 40) -> tuple[list[str], list[list[Any]], dict[str, Any]]:
    """Riadok hlavičky vo výkaze: prvý riadok, ktorý má aspoň 60 % šírky tabuľky vyplnenej a väčšinou text;
    riadky nad ním (nadpis, poznámky) sa vynechajú. Ďalší textový riadok pod hlavičkou, po ktorom idú dáta,
    sa spojí s hlavičkou (viacriadková hlavička; zlúčené bunky sa doplnia zľava)."""
    rows = [[("" if v is None else v) for v in r] for r in rows]
    for r in rows:
        while r and str(r[-1]).strip() == "":
            r.pop()
    filled = [sum(1 for v in r if str(v).strip() != "") for r in rows[:scan]]
    width = max(filled or [0])
    if width == 0:
        raise DatasetError("tabuľka je prázdna")
    header_i = None
    for i, r in enumerate(rows[:scan]):
        cells = [v for v in r if str(v).strip() != ""]
        if len(cells) >= max(2, 0.6 * width) and sum(1 for v in cells if _is_texty(v)) >= 0.6 * len(cells):
            header_i = i
            break
    info: dict[str, Any] = {"riadok_hlavičky": None, "vynechané_riadky_nad_tabuľkou": 0, "viacriadková_hlavička": False}
    if header_i is None:
        names = [f"column{j}" for j in range(width)]
        data = [r for r in rows if any(str(v).strip() for v in r)]
        return names, data, info
    header = list(rows[header_i])
    start = header_i + 1
    nxt = rows[start] if start < len(rows) else []
    nxt_cells = [v for v in nxt if str(v).strip() != ""]
    after = rows[start + 1] if start + 1 < len(rows) else []
    span = header + [""] * (width - len(header))
    gaps = [j for j, v in enumerate(span) if str(v).strip() == ""]  # zlúčené bunky hornej hlavičky
    fills_gaps = any(j < len(nxt) and str(nxt[j]).strip() for j in gaps)
    if (gaps and fills_gaps and nxt_cells and len(nxt_cells) >= 0.5 * width
            and sum(1 for v in nxt_cells if _is_texty(v) or _is_year(v)) >= 0.6 * len(nxt_cells)
            and any(_is_number(v) for v in after)):
        upper, last = [], ""
        for v in header + [""] * (len(nxt) - len(header)):
            last = str(v).strip() or last
            upper.append(last)
        header = [" ".join(x for x in (str(u).strip(), str(lw).strip()) if x) for u, lw in
                  zip(upper, list(nxt) + [""] * (len(upper) - len(nxt)))]
        start += 1
        info["viacriadková_hlavička"] = True
    names, seen = [], Counter()
    for j, h in enumerate(header + [""] * (width - len(header))):
        name = " ".join(str(h).split()) or f"stĺpec_{j + 1}"
        seen[name] += 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    data = []
    for r in rows[start:]:
        cells = [v for v in r if str(v).strip() != ""]
        if not cells:
            continue
        if len(cells) == 1 and width >= 3 and not _is_number(cells[0]):
            continue  # poznámka pod tabuľkou („Zdroj: ŠÚ SR“)
        data.append(list(r[: len(names)]) + [""] * (len(names) - len(r)))
    info.update(riadok_hlavičky=header_i + 1, vynechané_riadky_nad_tabuľkou=header_i)
    return names, data, info


def _write_csv(path: Path, names: list[str], data: list[list[Any]]) -> Path:
    import csv

    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(names)
        for r in data:
            w.writerow(["" if v is None else v for v in r])
    return path


def spreadsheetml_to_csv(path: Path) -> tuple[Path, dict[str, Any]]:
    """Excel 2003 XML (SpreadsheetML): Workbook → Worksheet → Table → Row → Cell → Data. Bunky sa skladajú podľa
    ss:Index a ss:MergeAcross; tabuľka a hlavička ako pri XLSX (audit: StyleID, MergeAcross a Formula ako stĺpce)."""
    from lxml import etree

    ns = "urn:schemas-microsoft-com:office:spreadsheet"
    tree = etree.parse(str(path), etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=True))
    best: tuple[int, list[str], list[list[Any]], dict[str, Any], str] | None = None
    for ws in tree.getroot().iter(f"{{{ns}}}Worksheet"):
        rows = []
        for row in ws.iter(f"{{{ns}}}Row"):
            cells: list[Any] = []
            for cell in row.findall(f"{{{ns}}}Cell"):
                idx = cell.get(f"{{{ns}}}Index")
                if idx and idx.isdigit():
                    cells += [""] * max(0, int(idx) - 1 - len(cells))
                data = cell.find(f"{{{ns}}}Data")
                cells.append("".join(data.itertext()).strip() if data is not None else "")
                cells += [""] * int(cell.get(f"{{{ns}}}MergeAcross") or 0)
            rows.append(cells)
        try:
            names, data, info = detect_table(rows)
        except DatasetError:
            continue
        if best is None or len(data) * len(names) > best[0]:
            best = (len(data) * len(names), names, data, info, ws.get(f"{{{ns}}}Name") or "")
    if best is None:
        raise DatasetError("SpreadsheetML bez tabuľky")
    _, names, data, info, sheet = best
    info["hárok"] = sheet
    return _write_csv(path.with_suffix(".ssml.csv"), names, data), info


# ---------------------------------------------------------------- RDF a JSON-LD (štatistické dátové kocky)
RDF_NS = "http://www.w3.org/1999/02/22-rdf-syntax-ns#"
LABEL_PROPS = ("prefLabel", "label", "notation", "name", "title", "hasValue", "value")
UNIT_PROPS = ("unitMeasure", "unit", "measureUnit", "jednotka", "unitOfMeasure")


def _uri_local(uri: str) -> str:
    return re.split(r"[#/]", uri.rstrip("/#"))[-1] if uri else ""


def _resolve(nodes: dict[str, dict[str, Any]], kind: str, value: str, depth: int = 0) -> str:
    """Hodnota vlastnosti pozorovania: odkaz na uzol → jeho označenie (prefLabel, label, notation, hasValue),
    URI bez uzla → posledná časť URI; literál ostáva."""
    if kind == "lit":
        return value
    node = nodes.get(value)
    if node is not None and depth < 3:
        for prop in LABEL_PROPS:
            for k, v in node["props"].get(prop, []):
                got = _resolve(nodes, k, v, depth + 1)
                if got:
                    return got
    return _uri_local(value) if kind == "uri" and not value.startswith("_:") else ""  # prázdny uzol bez označenia


TIME_DIM_RE = re.compile(r"(rok|year|period|obdob|time|čas|mesiac|month|kvart|quarter|date|dátum)", re.I)


def _cube_records(nodes: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], set[str]]:
    """Pozorovania a vlastnosti, ktoré sú odkazom na iný uzol alebo číselník (dimenzie, nie hodnoty)."""
    out, linked = [], set()
    for nid, node in nodes.items():
        if not (any(t.endswith("Observation") for t in node["types"]) or "dataSet" in node["props"]):
            continue
        rec: dict[str, Any] = {}
        for prop, vals in node["props"].items():
            if prop in ("type", "dataSet"):
                continue
            kind, value = vals[0]
            if kind != "lit":
                linked.add(prop)
            rec[prop] = _resolve(nodes, kind, value)
        out.append(rec)
    return out, linked


def rdf_nodes(path: Path) -> dict[str, dict[str, Any]]:
    """RDF/XML → uzly {id: {types, props: {vlastnosť: [(lit|uri|ref, hodnota)]}}}; odkazy rdf:nodeID a rdf:resource
    sa zachovajú, aby sa dimenzie pozorovania dali pripojiť (audit: odpojené nodeID/hasValue)."""
    from lxml import etree

    root = etree.parse(str(path), etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False,
                                                  huge_tree=True)).getroot()
    nodes: dict[str, dict[str, Any]] = {}
    counter = [0]

    def walk(el) -> str:
        nid = el.get(f"{{{RDF_NS}}}about") or (f"_:{el.get(f'{{{RDF_NS}}}nodeID')}" if el.get(f"{{{RDF_NS}}}nodeID") else None)
        if nid is None:
            counter[0] += 1
            nid = f"_:anon{counter[0]}"
        node = nodes.setdefault(nid, {"types": [], "props": {}})
        if _local(el.tag) != "Description":
            node["types"].append(_local(el.tag))
        for prop in el:
            name = _local(prop.tag)
            if not name:
                continue
            res, ref = prop.get(f"{{{RDF_NS}}}resource"), prop.get(f"{{{RDF_NS}}}nodeID")
            if name == "type" and res:
                node["types"].append(_uri_local(res))
                continue
            if res:
                val = ("uri", res)
            elif ref:
                val = ("ref", f"_:{ref}")
            elif len(prop):
                val = ("ref", walk(prop[0]))
            else:
                val = ("lit", (prop.text or "").strip())
            node["props"].setdefault(name, []).append(val)
        return nid

    for el in root:
        walk(el)
    for node in nodes.values():  # „ref“ a „uri“ na uzol v dokumente sa riešia rovnako
        for vals in node["props"].values():
            for i, (k, v) in enumerate(vals):
                if k == "ref":
                    vals[i] = ("uri", v)
    return nodes


def jsonld_nodes(path: Path) -> dict[str, dict[str, Any]]:
    import json

    data = json.loads(path.read_text(encoding="utf-8-sig"))
    items = data.get("@graph", []) if isinstance(data, dict) else data
    if isinstance(data, dict) and not items and "@id" in data:
        items = [data]
    nodes: dict[str, dict[str, Any]] = {}
    for i, item in enumerate(items or []):
        if not isinstance(item, dict):
            continue
        nid = str(item.get("@id") or f"_:j{i}")
        types = item.get("@type") or []
        node = nodes.setdefault(nid, {"types": [_uri_local(str(t)) for t in (types if isinstance(types, list) else [types])],
                                      "props": {}})
        for key, raw in item.items():
            if key.startswith("@"):
                continue
            name = _uri_local(key.split(":")[-1]) or key
            for v in (raw if isinstance(raw, list) else [raw]):
                if isinstance(v, dict):
                    val = ("uri", str(v["@id"])) if "@id" in v else ("lit", str(v.get("@value", "")))
                else:
                    val = ("lit", str(v))
                node["props"].setdefault(name, []).append(val)
    return nodes


def cube_to_jsonl(path: Path, kind: str) -> tuple[Path, dict[str, Any]]:
    """Pozorovania dátovej kocky (qb:Observation) s pripojenými dimenziami a jednotkou → JSONL. Bez pozorovaní,
    bez číselnej hodnoty alebo s prázdnymi dimenziami → chyba (neznámy formát / odpojené dimenzie)."""
    import json

    nodes = rdf_nodes(path) if kind == "rdf" else jsonld_nodes(path)
    records, linked = _cube_records(nodes)
    if not records:
        raise DatasetError("RDF/JSON-LD bez pozorovaní dátovej kocky (qb:Observation) – neznámy formát, na kontrolu")
    keys = sorted({k for r in records for k in r})
    dims = [k for k in keys if k not in UNIT_PROPS and (k in linked or TIME_DIM_RE.search(k))]
    numeric = [k for k in keys if k not in dims and k not in UNIT_PROPS
               and sum(1 for r in records if r.get(k) and _is_number(r[k])) >= 0.9 * len(records)]
    dims += [k for k in keys if k not in dims and k not in numeric and k not in UNIT_PROPS]
    if not numeric:
        raise DatasetError("pozorovania nemajú číselnú hodnotu – neznámy formát, na kontrolu")
    empty = sum(1 for r in records if not any(str(r.get(d) or "").strip() for d in dims))
    if not dims or empty > 0.1 * len(records):
        raise DatasetError(f"odpojené dimenzie: {empty} z {len(records)} pozorovaní bez roka, územia či ukazovateľa")
    out = path.with_suffix(".cube.jsonl")
    with open(out, "w", encoding="utf-8") as fh:
        for r in records:
            fh.write(json.dumps({k: r.get(k, "") for k in keys}, ensure_ascii=False) + "\n")
    unit = next((u for u in UNIT_PROPS if u in keys), None)
    values = sorted({str(r.get(unit) or "") for r in records} - {""})[:5] if unit else []
    return out, {"pozorovania": len(records), "dimenzie": dims, "hodnoty": numeric, "jednotka": unit,
                 "jednotky_hodnoty": values}


def normalize_csv(path: Path) -> tuple[Path, dict[str, Any]]:
    """CSV s nadpisom alebo poznámkami nad tabuľkou (exporty DATAcube, výkazy): hlavička sa nájde ako pri XLSX."""
    import csv

    with open(path, encoding="utf-8", errors="replace", newline="") as fh:
        sample = fh.read(16384)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        return path, {}
    with open(path, encoding="utf-8", errors="replace", newline="") as fh:
        head = [r for _, r in zip(range(40), csv.reader(fh, dialect))]
    try:
        names, _, info = detect_table(head)
    except DatasetError:
        return path, {}
    if not info.get("riadok_hlavičky") or (info["riadok_hlavičky"] == 1 and not info["viacriadková_hlavička"]):
        return path, info
    with open(path, encoding="utf-8", errors="replace", newline="") as fh:
        rows = list(csv.reader(fh, dialect))
    names, data, info = detect_table(rows, scan=40)
    return _write_csv(path.with_suffix(".table.csv"), names, data), info


def unpack_zip(cfg: Config, path: Path) -> tuple[Path, str, str]:
    """Najväčší podporovaný súbor z archívu (CSV, XLSX, JSON, Parquet, XML) → (cesta, formát, meno v archíve).
    Bezpečne: meno súboru z archívu sa nepoužije ako cesta (../), veľkosť je obmedzená (ochrana pred „zip bombou“)."""
    import zipfile

    limit = cfg.gate1.max_dataset_mb * 1024 * 1024
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise DatasetError("poškodený ZIP") from exc
    with zf:
        members = [m for m in zf.infolist() if not m.is_dir() and m.filename.rsplit(".", 1)[-1].lower() in TABLE_FORMATS]
        if not members:
            raise DatasetError("ZIP neobsahuje CSV, XLSX, JSON, Parquet ani XML")
        member = max(members, key=lambda m: m.file_size)
        if member.file_size > limit:
            raise DatasetError(f"súbor v ZIP má {member.file_size / 1e6:.0f} MB, limit {cfg.gate1.max_dataset_mb} MB")
        ext = member.filename.rsplit(".", 1)[-1].lower()
        fmt = "csv" if ext in ("tsv", "txt") else ext
        out = path.with_suffix(f".unzipped.{fmt}")
        written = 0
        with zf.open(member) as src, open(out, "wb") as dst:
            while chunk := src.read(1 << 20):
                written += len(chunk)
                if written > limit:  # deklarovaná veľkosť môže klamať
                    dst.close()
                    out.unlink(missing_ok=True)
                    raise DatasetError(f"súbor v ZIP presiahol limit {cfg.gate1.max_dataset_mb} MB")
                dst.write(chunk)
    return out, fmt, member.filename


def _local(tag: Any) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def xml_to_jsonl(path: Path, sample_events: int = 20000) -> Path:
    """XML so zoznamom záznamov → JSONL. Záznam = najčastejší prvok s podprvkami (napr. <projekt>, <zakazka>);
    stĺpce = atribúty a podprvky (jedna úroveň vnorenia ako „rodič_dieťa“). Streamovane, bez externých entít (XXE)."""
    import json
    from collections import Counter

    from lxml import etree

    counts: Counter = Counter()
    for n, (_, el) in enumerate(etree.iterparse(str(path), events=("end",), resolve_entities=False, no_network=True,
                                                load_dtd=False)):
        if len(el):
            counts[_local(el.tag)] += 1
        if n >= sample_events:
            break
    root_like = {t for t, c in counts.items() if c == 1}
    candidates = [(c, t) for t, c in counts.items() if c > 1 and t not in root_like]
    if not candidates:
        raise DatasetError("XML neobsahuje opakujúce sa záznamy")
    record_tag = max(candidates)[1]
    out = path.with_suffix(".xml.jsonl")
    rows = 0
    with open(out, "w", encoding="utf-8") as fh:
        for _, el in etree.iterparse(str(path), events=("end",), resolve_entities=False, no_network=True, load_dtd=False):
            if _local(el.tag) != record_tag:
                continue
            rec: dict[str, Any] = {_local(k): v for k, v in el.attrib.items()}
            for child in el:
                name = _local(child.tag)
                if not name:
                    continue
                if len(child):
                    for sub in child:
                        if _local(sub.tag):
                            rec.setdefault(f"{name}_{_local(sub.tag)}", (sub.text or "").strip())
                else:
                    rec.setdefault(name, (child.text or "").strip())
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
            rows += 1
            el.clear()
    if not rows:
        raise DatasetError("XML neobsahuje záznamy")
    return out


def download(cfg: Config, fetcher, url: str, dest_dir: Path, blocked: list[str]) -> Download:
    """Stiahne súbor po častiach; preruší pri prekročení limitu. Presmerovanie na blokovanú doménu = chyba."""
    limit = cfg.gate1.max_dataset_mb * 1024 * 1024
    if domain_blocked(url, blocked):
        raise DatasetError("blokovaná doména")
    if not fetcher.allowed(url):
        raise DatasetError("robots.txt nepovoľuje")
    dest_dir.mkdir(parents=True, exist_ok=True)
    tmp = dest_dir / ("download-" + hashlib.sha1(url.encode()).hexdigest()[:12] + ".part")
    digest = hashlib.sha256()
    size = 0
    try:
        with fetcher.http.stream("GET", url, follow_redirects=True) as resp:
            hops = [str(r.url) for r in resp.history] + [str(resp.url)]
            if any(domain_blocked(h, blocked) for h in hops):
                raise DatasetError("presmerovanie na blokovanú doménu")
            if resp.status_code >= 400:
                raise DatasetError(f"HTTP {resp.status_code}")
            declared = int(resp.headers.get("content-length") or 0)
            if declared > limit:
                raise DatasetError(f"súbor má {declared / 1e6:.0f} MB, limit {cfg.gate1.max_dataset_mb} MB")
            with open(tmp, "wb") as out:
                for chunk in resp.iter_bytes(1 << 20):
                    size += len(chunk)
                    if size > limit:
                        raise DatasetError(f"súbor presiahol limit {cfg.gate1.max_dataset_mb} MB")
                    digest.update(chunk)
                    out.write(chunk)
            ctype = resp.headers.get("content-type", "")
            final = str(resp.url)
    except DatasetError:
        tmp.unlink(missing_ok=True)
        raise
    except Exception as exc:  # noqa: BLE001 – sieťová chyba: dataset sa zapíše ako chyba, beh pokračuje
        tmp.unlink(missing_ok=True)
        raise DatasetError(f"sťahovanie zlyhalo: {exc}") from exc
    sha = digest.hexdigest()
    fmt = file_format(final, ctype) or file_format(url, ctype) or "bin"
    path = dest_dir / f"{sha[:16]}.{fmt}"
    tmp.replace(path)
    return Download(path, sha, size, final, ctype)


LIST_KEYS = ("data", "items", "results", "content", "records", "rows", "value", "hodnoty", "zoznam")


def _records(data: Any) -> list[Any] | None:
    """Záznamy z odpovede API: zoznam, alebo zoznam pod bežným kľúčom (data, items, results …)."""
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for key in LIST_KEYS:
            if isinstance(data.get(key), list):
                return data[key]
    return None


def download_paginated(cfg: Config, fetcher, url: str, dest_dir: Path, blocked: list[str],
                       paging: dict[str, Any]) -> Download:
    """API po stranách → JSONL (jeden záznam na riadok). Kurzor (napr. ITMS minId = posledné id), offset alebo
    číslo strany; koniec pri prázdnej alebo neúplnej strane, opakovanej strane, limite strán alebo veľkosti."""
    import json
    from urllib.parse import parse_qsl, urlencode, urlparse

    limit_bytes = cfg.gate1.max_dataset_mb * 1024 * 1024
    size_param = paging.get("limit")
    page_size = cfg.gate1.api_page_size
    parsed = urlparse(url)
    base_q = dict(parse_qsl(parsed.query))
    if size_param:
        base_q[size_param] = str(page_size)
    dest_dir.mkdir(parents=True, exist_ok=True)
    tmp = dest_dir / ("api-" + hashlib.sha1(url.encode()).hexdigest()[:12] + ".part")
    digest = hashlib.sha256()
    size = rows = pages = 0
    cursor, offset, page_no = None, 0, int(paging.get("first_page", 1))
    seen: set[str] = set()
    note = None
    with open(tmp, "wb") as out:
        for _ in range(cfg.gate1.max_api_pages):
            q = dict(base_q)
            if paging.get("cursor") and cursor is not None:
                q[paging["cursor"]] = str(cursor)
            elif paging.get("offset"):
                q[paging["offset"]] = str(offset)
            elif paging.get("page"):
                q[paging["page"]] = str(page_no)
            page_url = parsed._replace(query=urlencode(q)).geturl()
            if domain_blocked(page_url, blocked):
                raise DatasetError("blokovaná doména")
            records = _records(fetcher.fetch_json(page_url))
            if records is None:
                break  # na prvej strane = chyba (nižšie), neskôr = koniec dát
            if not records:
                break
            first = json.dumps(records[0], sort_keys=True, ensure_ascii=False)[:300]
            if first in seen:  # server ignoruje stránkovanie a vracia stále to isté
                break
            seen.add(first)
            for rec in records:
                line = (json.dumps(rec, ensure_ascii=False) + "\n").encode("utf-8")
                size += len(line)
                digest.update(line)
                out.write(line)
            rows += len(records)
            pages += 1
            if size > limit_bytes:
                note = f"skrátené na limit {cfg.gate1.max_dataset_mb} MB"
                break
            if not size_param or len(records) < page_size or not any(paging.get(k) for k in ("cursor", "offset", "page")):
                break
            if paging.get("cursor"):
                ids = [r.get("id") for r in records if isinstance(r, dict) and isinstance(r.get("id"), (int, float))]
                if not ids:
                    break
                cursor = max(ids)
            offset += len(records)
            page_no += 1
        else:
            note = f"skrátené na {cfg.gate1.max_api_pages} strán"
    if rows == 0:
        tmp.unlink(missing_ok=True)
        raise DatasetError("API nevrátilo zoznam záznamov")
    sha = digest.hexdigest()
    path = dest_dir / f"{sha[:16]}.jsonl"
    tmp.replace(path)
    return Download(path, sha, size, url, "application/x-ndjson",
                    f"API: {pages} strán, {rows} záznamov" + (f"; {note}" if note else ""))


def _to_utf8(path: Path) -> Path:
    """CSV z úradov býva v cp1250; DuckDB číta UTF-8. Prekódovanie po riadkoch (bez načítania celého súboru)."""
    with open(path, "rb") as fh:
        head = fh.read(1 << 20)
    try:
        head.decode("utf-8")
        return path
    except UnicodeDecodeError:
        pass
    out = path.with_suffix(".utf8.csv")
    with open(path, encoding="cp1250", errors="replace", newline="") as src, open(out, "w", encoding="utf-8", newline="") as dst:
        for line in src:
            dst.write(line)
    return out


def _xlsx_to_csv(path: Path) -> tuple[Path, dict[str, Any]]:
    """Hárok s najväčšou tabuľkou (riadky × stĺpce po nájdení hlavičky), nie s najviac riadkami; nadpis a poznámky
    nad tabuľkou sa vynechajú, viacriadková hlavička sa spojí (audit: hlavičky „Filter na pozadí“, „column02“)."""
    from openpyxl import load_workbook

    with open(path, "rb") as fh:  # súbor ako objekt: openpyxl inak odmietne príponu (stiahnuté .bin/.part)
        wb = load_workbook(BytesIO(fh.read()), read_only=True, data_only=True)
    best: tuple[int, list[str], list[list[Any]], dict[str, Any]] | None = None
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        try:
            names, data, info = detect_table(rows)
        except DatasetError:
            continue
        info["hárok"] = ws.title
        size = len(data) * len(names)
        if best is None or size > best[0]:
            best = (size, names, data, info)
    wb.close()
    if best is None:
        raise DatasetError("XLSX bez tabuľky")
    _, names, data, info = best
    return _write_csv(path.with_suffix(".csv"), names, data), info


def _quote(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def import_table(cfg: Config, hid: int, table: str, path: Path, fmt: str) -> dict[str, Any]:
    """Súbor → tabuľka v DuckDB databáze hypotézy (zápis; pred dotazmi modelu sa otvorí iba na čítanie).
    Formát podľa obsahu (HTML sa odmietne); ZIP sa rozbalí (najväčší podporovaný súbor); SpreadsheetML a XLSX na CSV
    s nájdenou hlavičkou; RDF a JSON-LD dátové kocky na pozorovania s dimenziami; ostatné XML na záznamy (overené
    kontrolou); CSV v cp1250 na UTF-8 a odmietnuté riadky sa spočítajú. Vráti informácie o importe (adaptér, verzia
    parsera, odmietnuté riadky, hlavička). Všetky pomocné súbory sa po importe (aj pri chybe) zmažú."""
    if not re.fullmatch(r"ds\d+", table):
        raise DatasetError("neplatný názov tabuľky")
    cleanup: list[Path] = []
    info: dict[str, Any] = {"verzia_parsera": PARSER_VERSION}
    try:
        src, fmt = path, sniff_format(path, fmt)
        if fmt == "zip":
            src, fmt, member = unpack_zip(cfg, path)
            cleanup.append(src)
            info["súbor_v_zip"] = member
            fmt = sniff_format(src, fmt)
        info["adaptér"] = fmt
        if fmt == "spreadsheetml":
            try:
                src, extra = spreadsheetml_to_csv(src)
            except DatasetError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise DatasetError(f"SpreadsheetML sa nedá spracovať: {str(exc)[:150]}") from exc
            info["tabuľka"], fmt = extra, "csv"
            cleanup.append(src)
        elif fmt in ("rdf", "jsonld"):
            try:
                src, extra = cube_to_jsonl(src, fmt)
            except DatasetError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise DatasetError(f"RDF/JSON-LD sa nedá spracovať: {str(exc)[:150]}") from exc
            info["kocka"], fmt = extra, "json"
            cleanup.append(src)
        elif fmt == "xml":
            try:
                src, fmt = xml_to_jsonl(src), "json"
            except DatasetError:
                raise
            except Exception as exc:  # noqa: BLE001 – poškodené XML: chyba datasetu, beh pokračuje
                raise DatasetError(f"XML sa nedá spracovať: {str(exc)[:150]}") from exc
            info["adaptér"] = "xml_záznamy"
            cleanup.append(src)
        if fmt == "xlsx":
            src, extra = _xlsx_to_csv(src)
            info["tabuľka"], fmt = extra, "csv"
            cleanup.append(src)
        if fmt == "csv":
            utf8 = _to_utf8(src)
            if utf8 != src:
                cleanup.append(utf8)
            src = utf8
            if "tabuľka" not in info:
                normalized, extra = normalize_csv(src)
                if normalized != src:
                    cleanup.append(normalized)
                    src = normalized
                if extra:
                    info["tabuľka"] = extra
            reader = "read_csv(?, auto_detect = true, ignore_errors = true, store_rejects = true, sample_size = 20000)"
        elif fmt == "json":
            # Odpovede API bývajú vnorené (prijimatel.subjekt.ico): rozbalia sa na stĺpce (ico, dic …).
            reader = "(SELECT unnest(j, recursive := true) FROM read_json_auto(?) AS j)"
        elif fmt == "parquet":
            reader = "read_parquet(?)"
        else:
            raise DatasetError(f"nepodporovaný formát {fmt}")
        con = duckdb.connect(str(duck_path(cfg, hid)))
        try:
            try:
                con.execute(f"CREATE OR REPLACE TABLE {table} AS SELECT * FROM {reader}", [str(src)])
            except duckdb.Error:
                if fmt != "json":
                    raise
                con.execute(f"CREATE OR REPLACE TABLE {table} AS SELECT * FROM read_json_auto(?)", [str(src)])
            if fmt == "csv":
                info["odmietnuté_riadky"] = _rejected_rows(con)
                # Hlavička z rokov („2016;2017…“) má rovnaký typ ako číselné dáta, takže ju DuckDB nespozná
                # (stĺpce column0…); keď prvý riadok vyzerá ako obdobia, načíta sa znova s hlavičkou.
                names = [r[0] for r in con.execute(f"DESCRIBE {table}").fetchall()]
                if names and all(re.fullmatch(r"column\d+", n) for n in names):
                    first = con.execute(f"SELECT * FROM {table} LIMIT 1").fetchone() or ()
                    if period_columns([str(v) for v in first if v is not None], cfg.gate1.min_wide_periods)[1]:
                        con.execute(f"CREATE OR REPLACE TABLE {table} AS SELECT * FROM "
                                    f"{reader.replace('auto_detect = true', 'auto_detect = true, header = true')}", [str(src)])
                        info["odmietnuté_riadky"] = _rejected_rows(con)
        except duckdb.Error as exc:
            raise DatasetError(f"import zlyhal: {str(exc)[:200]}") from exc
        finally:
            con.close()
    except DatasetError:
        raise
    except Exception as exc:  # noqa: BLE001 – poškodený súbor (napr. XLSX) je chyba datasetu, nie celého kroku
        raise DatasetError(f"súbor sa nedá spracovať: {type(exc).__name__}: {str(exc)[:150]}") from exc
    finally:
        for extra_path in cleanup:
            if extra_path != path:
                extra_path.unlink(missing_ok=True)
    return info


def _rejected_rows(con) -> int:
    """Riadky CSV, ktoré DuckDB pri ignore_errors vynechal (zlé typy, počet stĺpcov) – podmienka kvality importu."""
    try:
        return int(con.execute("SELECT COUNT(DISTINCT (scan_id, line)) FROM reject_errors").fetchone()[0])
    except duckdb.Error:
        return 0


def assess_import(cfg: Config, hid: int, table: str, prof: dict[str, Any], info: dict[str, Any]) -> dict[str, Any]:
    """Importný report: načítané a odmietnuté riadky, stĺpce a typy, kľúče, jednotky, vzorka a problémy. Verdikt
    „ok“ iba bez problémov; inak „needs_review“ – dataset sa do pripravenosti ani experimentov neráta, kým ho
    človek neschváli (audit: „ok“ nezaručuje použiteľné dáta)."""
    import json

    names = [c["name"] for c in prof["columns"]]
    rows = int(prof.get("rows") or 0)
    rejected = int(info.get("odmietnuté_riadky") or 0)
    problems = []
    if rows == 0:
        problems.append("tabuľka je prázdna")
    share = rejected / (rows + rejected) if rows + rejected else 0.0
    if share > cfg.gate1.max_rejected_share:
        problems.append(f"parser odmietol {rejected} riadkov ({share:.0%})")
    unknown = [n for n in names if re.fullmatch(r"(column\d+|unnamed.*|stĺpec_\d+|)", n.strip().lower())]
    if names and len(unknown) > 0.3 * len(names):
        problems.append(f"neznáme hlavičky: {len(unknown)} z {len(names)} ({', '.join(unknown[:4])})")
    if any("<" in n or "doctype" in n.lower() for n in names):
        problems.append("hlavička obsahuje HTML alebo XML")
    fmt_cols = [n for n in names if n.strip().lower() in FORMAT_COLUMNS]
    if len(fmt_cols) >= 2 or (names and len(fmt_cols) >= 0.5 * len(names)):
        problems.append("stĺpce formátovania buniek alebo grafu (" + ", ".join(fmt_cols[:4]) + ") – nie záznamy s významom")
    texty_header = [n for n in names if _is_number(n) or re.fullmatch(r"\d{1,4}[./-]\d{1,2}([./-]\d{1,4})?", n.strip())]
    if names and len(texty_header) > 0.5 * len(names) and not prof.get("period_columns"):
        problems.append("hlavička vyzerá ako dáta (čísla, dátumy)")
    if info.get("adaptér") == "xml_záznamy" and len([n for n in names if n.strip().lower() not in FORMAT_COLUMNS]) < 3:
        problems.append("neznámy formát XML (záznamy s menej ako 3 poľami) – skontroluj vzorku")
    if info.get("adaptér") == "jsonld":
        pass  # dátová kocka s pozorovaniami (adaptér overil dimenzie a hodnoty)
    sample: list[dict[str, Any]] = []
    keys: list[str] = list(prof.get("ico_columns") or [])
    sums: list[float] = []
    if rows:
        con = connect_ro(duck_path(cfg, hid))
        try:
            cur = con.execute(f"SELECT * FROM {table} LIMIT 5")
            cols = [d[0] for d in cur.description]
            sample = [{c: (str(v)[:60] if v is not None else None) for c, v in zip(cols, r)} for r in cur.fetchall()]
            for c in prof["columns"][:60]:
                qn = _quote(c["name"])
                if c["type"] == "VARCHAR" and re.search(r"(^|_)(id|kod|kód|číslo|cislo)($|_)", c["name"], re.I):
                    distinct = con.execute(f"SELECT COUNT(DISTINCT {qn}) FROM {table}").fetchone()[0]
                    if distinct == rows:
                        keys.append(c["name"])
                if any(t in c["type"] for t in ("INT", "DOUBLE", "DECIMAL", "FLOAT")):
                    total = con.execute(f"SELECT SUM(TRY_CAST({qn} AS DOUBLE)) FROM {table}").fetchone()[0]
                    sums.append(round(float(total or 0.0), 2))
            if len(names) == 1 and sample and "<" in str(sample[0].get(names[0]) or ""):
                problems.append("XML/RDF v jednom textovom stĺpci – nespracovaný register")
        finally:
            con.close()
    units = {n: m.group(1) for n in names if (m := UNIT_RE.search(n))}
    cube = info.get("kocka") or {}
    if cube.get("jednotka"):
        units[cube["jednotka"]] = ", ".join(cube.get("jednotky_hodnoty") or []) or "chýba"
    fingerprint = hashlib.sha1(json.dumps([rows, sorted(sums)]).encode()).hexdigest()[:16] if sums else None
    return {**info, "riadky": rows, "odmietnuté_riadky": rejected, "stĺpce": len(names),
            "typy": {c["name"]: c["type"] for c in prof["columns"][:60]}, "kľúče": list(dict.fromkeys(keys)),
            "jednotky": units, "vzorka": sample, "problémy": problems, "odtlačok_obsahu": fingerprint,
            "verdikt": "needs_review" if problems else "ok"}


def raw_cache_dir(cfg: Config) -> Path:
    return cfg.path(cfg.gate1.datasets_dir) / "raw"


def raw_cache_put(cfg: Config, path: Path, sha256: str) -> Path | None:
    """Pôvodný súbor podľa SHA-256 (zopakovanie parsera po oprave); najstaršie sa mažú nad rozpočtom raw_cache_gb."""
    import shutil

    if not path.exists() or path.stat().st_size > cfg.gate1.raw_cache_max_file_mb * 1024 * 1024 or not sha256:
        return None
    folder = raw_cache_dir(cfg)
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{sha256}{path.suffix}"
    if not target.exists():
        shutil.copyfile(path, target)
    files = sorted(folder.iterdir(), key=lambda p: p.stat().st_mtime)
    total = sum(p.stat().st_size for p in files)
    budget = cfg.gate1.raw_cache_gb * 1024 ** 3
    for old in files:
        if total <= budget or old == target:
            continue
        total -= old.stat().st_size
        old.unlink(missing_ok=True)
    return target


def raw_cache_drop(cfg: Config, sha256: str) -> None:
    """Dataset s osobnými údajmi: pôvodný súbor sa neuchováva (minimalizácia)."""
    folder = raw_cache_dir(cfg)
    if sha256 and folder.exists():
        for p in folder.glob(f"{sha256}.*"):
            p.unlink(missing_ok=True)


def drop_columns(cfg: Config, hid: int, table: str, columns: list[str]) -> list[str]:
    """Odstráni stĺpce z tabuľky datasetu (minimalizácia osobných údajov po rozhodnutí človeka); vráti odstránené."""
    if not re.fullmatch(r"ds\d+", table) or not duck_path(cfg, hid).exists():
        return []
    con = duckdb.connect(str(duck_path(cfg, hid)))
    try:
        existing = {r[0] for r in con.execute(f"DESCRIBE {table}").fetchall()}
        dropped = [c for c in columns if c in existing]
        for c in dropped:
            con.execute(f"ALTER TABLE {table} DROP COLUMN {_quote(c)}")
        return dropped
    finally:
        con.close()


def drop_table(cfg: Config, hid: int, table: str) -> None:
    if not re.fullmatch(r"ds\d+", table) or not duck_path(cfg, hid).exists():
        return
    con = duckdb.connect(str(duck_path(cfg, hid)))
    try:
        con.execute(f"DROP TABLE IF EXISTS {table}")
        con.execute(f"DROP TABLE IF EXISTS {table}_long")
    finally:
        con.close()


def profile_table(cfg: Config, hid: int, table: str) -> dict[str, Any]:
    """Riadky, stĺpce (typ, podiel prázdnych, rozsah), časový rozsah a podozrenie na osobné údaje."""
    con = connect_ro(duck_path(cfg, hid))
    try:
        cols = con.execute(f"DESCRIBE {table}").fetchall()
        rows = con.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        columns, personal, time_cols, ico_cols = [], [], [], []
        gran, periods = period_columns([c[0] for c in cols], cfg.gate1.min_wide_periods)
        for name, ctype, *_ in cols:
            q = _quote(name)
            nulls = con.execute(f"SELECT COUNT(*) FILTER (WHERE {q} IS NULL) FROM {table}").fetchone()[0]
            info: dict[str, Any] = {"name": name, "type": ctype, "nulls": round(nulls / rows, 3) if rows else None}
            if any(t in ctype for t in ("INT", "DOUBLE", "DECIMAL", "FLOAT", "DATE", "TIMESTAMP", "BIGINT")):
                lo, hi = con.execute(f"SELECT MIN({q}), MAX({q}) FROM {table}").fetchone()
                info.update(min=str(lo), max=str(hi))
            if PERSONAL_COLUMN_RE.search(name):
                personal.append(name)
            elif "VARCHAR" in ctype:
                sample = [r[0] for r in con.execute(f"SELECT {q} FROM {table} WHERE {q} IS NOT NULL LIMIT 200").fetchall()]
                hits = sum(1 for v in sample if EMAIL_VALUE_RE.search(str(v)) or _is_phone(str(v)))
                if sample and hits / len(sample) > 0.2:
                    personal.append(name)
            if (TIME_COLUMN_RE.search(name) or "DATE" in ctype or "TIMESTAMP" in ctype) and name not in periods:
                time_cols.append(name)
            if ICO_COLUMN_RE.search(name):
                ico_cols.append(name)
            columns.append(info)
        span = {}
        for name in time_cols[:3]:
            info = next(c for c in columns if c["name"] == name)
            if "min" in info:
                span[name] = [info["min"], info["max"]]
        if periods:
            span["obdobia v stĺpcoch"] = [min(periods.values()).isoformat(), max(periods.values()).isoformat()]
        ids = {name: identifier_quality(con, table, name) for name in ico_cols[:6]}
        return {"rows": rows, "columns": columns, "time_columns": time_cols, "time_span": span,
                "ico_columns": ico_cols, "personal_columns": personal, "identifikátory": ids,
                "period_columns": {"granularita": gran, "počet": len(periods)} if periods else None}
    finally:
        con.close()


def period_of(name: str) -> tuple[str, date] | None:
    """(granularita, začiatok obdobia) z názvu stĺpca, napr. „2021-Q2“ → (quarter, 2021-04-01)."""
    s = str(name).strip()
    for gran, rx in PERIOD_NAME_RES:
        m = rx.match(s)
        if m:
            g = m.groupdict()
            month = int(g["m"]) if g.get("m") else (int(g["q"]) - 1) * 3 + 1 if g.get("q") else 1
            return gran, date(int(g["y"]), month, 1)
    return None


def period_columns(columns: list[str], min_periods: int = 3) -> tuple[str, dict[str, date]]:
    """Stĺpce, ktorých názov je obdobie (široká tabuľka): najčastejšia granularita, aspoň `min_periods` stĺpcov."""
    found: dict[str, dict[str, date]] = {}
    for c in columns:
        p = period_of(c)
        if p:
            found.setdefault(p[0], {})[c] = p[1]
    if not found:
        return "", {}
    gran, cols = max(found.items(), key=lambda kv: len(kv[1]))
    if len(cols) < min_periods:
        return "", {}
    return gran, cols


# Hodnota bunky → DOUBLE: „1 234,5“, „1.234,5“, „1,234.5“, „12 p“ (príznak Eurostatu); „:“, „-“ a prázdne → NULL.
# Pri jednej čiarke ide o desatinnú čiarku (SK/CZ), pri viacerých o oddeľovače tisícov.
NUMBER_MACRO = r"""CREATE OR REPLACE TEMP MACRO radar_num(x) AS (
  WITH s AS (SELECT regexp_replace(CAST(x AS VARCHAR), '[^0-9,.\-]', '', 'g') AS v)
  SELECT TRY_CAST(CASE
    WHEN contains(v, ',') AND contains(v, '.') THEN
      CASE WHEN strpos(v, ',') > strpos(v, '.') THEN replace(replace(v, '.', ''), ',', '.') ELSE replace(v, ',', '') END
    WHEN length(v) - length(replace(v, ',', '')) > 1 THEN replace(v, ',', '')
    WHEN contains(v, ',') THEN replace(v, ',', '.')
    WHEN length(v) - length(replace(v, '.', '')) > 1 THEN replace(v, '.', '')
    ELSE v END AS DOUBLE) FROM s)"""


def wide_to_long(cfg: Config, hid: int, table: str) -> dict[str, Any] | None:
    """Široká tabuľka (obdobia ako stĺpce) → `{table}_long`: identifikačné stĺpce, obdobie (DATE), obdobie_stlpec
    (pôvodný názov stĺpca – pri viacerých ukazovateľoch na obdobie sa podľa neho filtruje) a hodnota (DOUBLE)."""
    if not re.fullmatch(r"ds\d+", table) or not duck_path(cfg, hid).exists():
        return None
    con = duckdb.connect(str(duck_path(cfg, hid)))
    try:
        cols = [r[0] for r in con.execute(f"DESCRIBE {table}").fetchall()]
        gran, periods = period_columns(cols, cfg.gate1.min_wide_periods)
        if not periods:
            return None
        rows = con.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        if rows * len(periods) > cfg.gate1.max_long_rows:
            return {"tabuľka": None, "granularita": gran, "obdobia": len(periods),
                    "poznámka": f"dlhý tvar by mal {rows * len(periods)} riadkov (limit {cfg.gate1.max_long_rows})"}
        ids = [c for c in cols if c not in periods]
        taken = set(ids)

        def free(name: str) -> str:
            while name in taken:
                name = "_" + name
            taken.add(name)
            return name

        n_t, n_txt, n_v = free("obdobie"), free("obdobie_stlpec"), free("hodnota")
        head = "".join(_quote(c) + ", " for c in ids)
        parts = [f"SELECT {head}DATE '{d.isoformat()}' AS {_quote(n_t)}, '{c.replace(chr(39), chr(39) * 2)}' AS "
                 f"{_quote(n_txt)}, radar_num({_quote(c)}) AS {_quote(n_v)} FROM {table}" for c, d in periods.items()]
        long = f"{table}_long"
        con.execute(NUMBER_MACRO)
        con.execute(f"CREATE OR REPLACE TABLE {long} AS " + " UNION ALL ".join(parts))
        n, filled = con.execute(f"SELECT COUNT(*), COUNT({_quote(n_v)}) FROM {long}").fetchone()
    except duckdb.Error as exc:
        log.warning("Dlhý tvar %s sa nepodaril: %s", table, exc)
        return {"tabuľka": None, "granularita": gran, "obdobia": len(periods), "poznámka": f"chyba: {str(exc)[:150]}"}
    finally:
        con.close()
    return {"tabuľka": long, "granularita": gran, "obdobia": len(periods), "riadky": n, "hodnoty": filled,
            "rozsah": [min(periods.values()).isoformat(), max(periods.values()).isoformat()],
            "stĺpce": {"obdobie": n_t, "obdobie_stlpec": n_txt, "hodnota": n_v}}


def base_table(table: str) -> str:
    """Tabuľka datasetu, z ktorej vznikol dlhý tvar (ds5_long → ds5)."""
    return re.sub(r"_long$", "", table or "")


def ico_check_sql(key: str) -> str:
    """SQL: kontrolná číslica 8-miestneho IČO (SR/ČR): (11 − Σ číslica × váha 8…2 mod 11) mod 10 = posledná číslica."""
    total = " + ".join(f"CAST(substr({key}, {i + 1}, 1) AS INTEGER) * {8 - i}" for i in range(7))
    return f"((11 - (({total}) % 11)) % 10) = CAST(substr({key}, 8, 1) AS INTEGER)"


def identifier_quality(con, table: str, column: str) -> dict[str, Any]:
    """Kvalita identifikátora v stĺpci (pôvodné hodnoty ostávajú v tabuľke; porovnáva sa normalizovaná podoba):
    čisté 8-miestne IČO, normalizované (medzery, text „IČO:“, 6–7 číslic), nerozpoznané (napr. zahraničné ID)
    a platná kontrolná číslica (audit: v CRZ 50 070 riadkov IČO objednávateľa s nečíselnými znakmi)."""
    col = _quote(column)
    key = ico_sql(col)
    filled, clean, norm, valid = con.execute(
        f"SELECT COUNT({col}), COUNT(*) FILTER (WHERE regexp_matches(CAST({col} AS VARCHAR), '^[0-9]{{8}}$')), "
        f"COUNT({key}), COUNT(*) FILTER (WHERE {key} IS NOT NULL AND {ico_check_sql(key)}) FROM {table}").fetchone()
    return {"vyplnené": int(filled), "čisté_8_číslic": int(clean), "normalizované": int(norm) - int(clean),
            "nerozpoznané": int(filled) - int(norm), "kontrolná_číslica_ok": int(valid)}


def ico_sql(expr: str) -> str:
    """SQL: IČO zo stĺpca ako 8 číslic – prvá samostatná skupina 6–8 číslic (aj „IČO: 00322521“ alebo „51 049 775“
    z CRZ), inak NULL."""
    text = f"replace(CAST({expr} AS VARCHAR), ' ', '')"
    return f"lpad(NULLIF(regexp_extract({text}, '(^|[^0-9])([0-9]{{6,8}})([^0-9]|$)', 2), ''), 8, '0')"


# Strana zmluvy, ktorá je samotná organizácia (objednávateľ, obstarávateľ, prijímateľ), nie jej protistrana (dodávateľ).
OWN_SIDE_RE = re.compile(r"objedn|obstar|zadav|prijim|ziadat|odberat|kupuj", re.I)


def _ascii(text: str) -> str:
    import unicodedata

    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()


def coverage(cfg: Config, hid: int, table: str, ico_columns: list[str], orgs: list[tuple[int, str]]) -> list[tuple[int, str, int]]:
    """(org_id, stĺpec, počet riadkov) pre organizácie z DB, ktorých IČO je v datasete – jeden stĺpec na organizáciu:
    prednostne jej vlastná strana (objednávateľ, prijímateľ), inak stĺpec s najviac riadkami. MF SR je v CRZ
    dodávateľom tisícok zmlúv obcí (výpomoci), no jeho vlastné výdavky sú tam, kde je objednávateľom."""
    by_ico = {str(ico).zfill(8): oid for oid, ico in orgs if ico}
    if not by_ico or not ico_columns:
        return []
    con = connect_ro(duck_path(cfg, hid))
    found: dict[int, list[tuple[str, int]]] = {}
    try:
        for col in ico_columns:
            ico = ico_sql(_quote(col))
            rows = con.execute(f"SELECT {ico} AS ico, COUNT(*) FROM {table} WHERE {ico} IN (SELECT unnest(?)) GROUP BY 1",
                               [list(by_ico)]).fetchall()
            for value, n in rows:
                found.setdefault(by_ico[value], []).append((col, n))
    finally:
        con.close()
    out = []
    for org_id, cols in found.items():
        own = [c for c in cols if OWN_SIDE_RE.search(_ascii(c[0]))]
        col, n = max(own or cols, key=lambda c: c[1])
        out.append((org_id, col, n))
    return out


def schema_text(cfg: Config, hid: int, tables: list[str], samples: int = 3) -> str:
    """Popis tabuliek pre model: stĺpce, typy a niekoľko riadkov (osobné údaje sa sem nedostanú – také datasety
    sa neukladajú)."""
    if not duck_path(cfg, hid).exists():
        return ""
    con = connect_ro(duck_path(cfg, hid))
    parts = []
    try:
        existing = {r[0] for r in con.execute("SELECT table_name FROM information_schema.tables").fetchall()}
        for t in [n for t in tables for n in (t, f"{t}_long") if n in existing]:
            cols = con.execute(f"DESCRIBE {t}").fetchall()
            if f"{t}_long" in existing:  # široká tabuľka: stĺpce období iba zhrnuté (stovky mesiacov by zahltili prompt)
                _, periods = period_columns([c[0] for c in cols], 1)
                n = con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
                head = ", ".join(f"{c[0]} {c[1]}" for c in cols if c[0] not in periods)
                parts.append(f"TABLE {t} ({n} rows, wide): {head}, plus {len(periods)} period columns "
                             f"{min(periods, key=periods.get)}…{max(periods, key=periods.get)} – use {t}_long for time series")
                continue
            rows = con.execute(f"SELECT * FROM {t} LIMIT {samples}").fetchall()
            n = con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
            head = ", ".join(f"{c[0]} {c[1]}" for c in cols)
            sample = "\n".join("  " + " | ".join(str(v)[:40] for v in r) for r in rows)
            values = _category_values(con, t, [c[0] for c in cols if c[1] == "VARCHAR"])
            parts.append(f"TABLE {t} ({n} rows): {head}\n{sample}" + (f"\n  values: {values}" if values else ""))
    finally:
        con.close()
    return "\n\n".join(parts)


def planner_profile(cfg: Config, hid: int, tables: list[str], max_columns: int = 40) -> str:
    """Profil tabuliek pre návrh experimentu: pre každý stĺpec typ, podiel prázdnych, približný počet rôznych hodnôt,
    rozsah čísel alebo dátumov, kódy s početnosťami alebo ukážky textu. Model z neho vyberá stĺpce a píše výrazy;
    bez neho hádal (napr. roky pri 9-mesačnom okne)."""
    if not duck_path(cfg, hid).exists():
        return ""
    con = connect_ro(duck_path(cfg, hid))
    parts = []
    try:
        existing = {r[0] for r in con.execute("SELECT table_name FROM information_schema.tables").fetchall()}
        for t in [n for t in tables for n in (f"{t}_long", t) if n in existing]:
            if f"{t}_long" in existing and not t.endswith("_long"):
                continue  # široká tabuľka: časové rady sú v dlhom tvare
            n = con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
            lines = [f"TABLE {t} ({n} rows)"]
            for name, ctype, *_ in con.execute(f"DESCRIBE {t}").fetchall()[:max_columns]:
                c = _quote(name)
                nulls, distinct = con.execute(f"SELECT AVG(CASE WHEN {c} IS NULL THEN 1.0 ELSE 0 END), "
                                              f"approx_count_distinct({c}) FROM {t}").fetchone()
                info = f"- {name} {ctype} | null {round((nulls or 0) * 100)}% | distinct ~{distinct}"
                if any(k in ctype for k in ("INT", "DOUBLE", "DECIMAL", "FLOAT", "BIGINT", "HUGEINT")):
                    lo, med, hi, zero = con.execute(
                        f"SELECT MIN({c}), approx_quantile({c}, 0.5), MAX({c}), AVG(CASE WHEN {c} = 0 THEN 1.0 ELSE 0 END) "
                        f"FROM {t}").fetchone()
                    info += f" | min {lo} median {med} max {hi} | zero {round((zero or 0) * 100)}%"
                elif "DATE" in ctype or "TIMESTAMP" in ctype:
                    lo, hi = con.execute(f"SELECT MIN({c}), MAX({c}) FROM {t}").fetchone()
                    info += f" | {lo} .. {hi}"
                elif distinct <= 12:
                    vals = con.execute(f"SELECT {c}, COUNT(*) FROM {t} GROUP BY 1 ORDER BY 2 DESC LIMIT 12").fetchall()
                    info += " | values: " + ", ".join(f"{str(v)[:30]} ({k})" for v, k in vals)
                else:
                    dateish = con.execute(f"SELECT AVG(CASE WHEN TRY_CAST({c} AS TIMESTAMP) IS NOT NULL THEN 1.0 ELSE 0 END) "
                                          f"FROM (SELECT {c} FROM {t} WHERE {c} IS NOT NULL LIMIT 20000)").fetchone()[0]
                    if (dateish or 0) > 0.9:
                        lo, hi = con.execute(f"SELECT MIN(TRY_CAST({c} AS TIMESTAMP)), MAX(TRY_CAST({c} AS TIMESTAMP)) "
                                             f"FROM {t}").fetchone()
                        info += f" | date-like text {lo} .. {hi}"
                    else:
                        ex = con.execute(f"SELECT {c}, AVG(length(CAST({c} AS VARCHAR))) OVER () FROM {t} "
                                         f"WHERE {c} IS NOT NULL LIMIT 3").fetchall()
                        info += " | text, e.g. " + " / ".join(repr(str(r[0])[:60]) for r in ex[:2])
                lines.append(info)
            parts.append("\n".join(lines))
    finally:
        con.close()
    return "\n\n".join(parts)


def _category_values(con, table: str, columns: list[str], max_values: int = 30, max_columns: int = 8) -> str:
    """Všetky hodnoty textových stĺpcov s malým počtom hodnôt (kraj, ukazovateľ): model z nich berie presné filtre.
    Počíta sa zo vzorky 100 000 riadkov (veľké tabuľky); stĺpce s mnohými hodnotami (názvy, IČO) sa vynechajú."""
    out = []
    for col in columns:
        if len(out) >= max_columns:
            break
        q = _quote(col)
        vals = con.execute(f"SELECT DISTINCT {q} FROM (SELECT {q} FROM {table} LIMIT 100000) WHERE {q} IS NOT NULL "
                           f"LIMIT {max_values + 1}").fetchall()
        if 2 <= len(vals) <= max_values:
            out.append(f"{col} = " + " | ".join(sorted(str(v[0])[:40] for v in vals)))
    return "; ".join(out)


def check_sql(sql: str) -> str:
    """Iba jeden SELECT (alebo WITH … SELECT) nad tabuľkami; žiadne súbory, sieť ani zmeny."""
    s = (sql or "").strip().rstrip(";").strip()
    if not re.match(r"^(select|with)\b", s, re.I):
        raise DatasetError("dotaz musí začínať SELECT alebo WITH")
    if FORBIDDEN_SQL.search(s):
        raise DatasetError("dotaz obsahuje zakázaný príkaz (súbory, sieť, zmeny alebo viac príkazov)")
    return s


def run_query(cfg: Config, hid: int, sql: str, params: list[Any] | None = None, timeout_s: float | None = None) -> list[tuple]:
    """Dotaz iba na čítanie, bez prístupu k súborom a sieti, s časovým limitom (con.interrupt)."""
    s = check_sql(sql)
    path = duck_path(cfg, hid)
    if not path.exists():
        raise DatasetError("hypotéza nemá žiadny dataset")
    con = connect_ro(path, {"enable_external_access": False})
    timer = threading.Timer(timeout_s or cfg.gate1.sql_timeout_s, con.interrupt)
    timer.start()
    try:
        return con.execute(s, params or []).fetchmany(1000)
    except duckdb.Error as exc:
        raise DatasetError(f"dotaz zlyhal: {str(exc)[:200]}") from exc
    finally:
        timer.cancel()
        con.close()


def table_columns(cfg: Config, hid: int, table: str) -> list[str]:
    if not TABLE_RE.fullmatch(table or "") or not duck_path(cfg, hid).exists():
        return []
    con = connect_ro(duck_path(cfg, hid))
    try:
        return [r[0] for r in con.execute(f"DESCRIBE {table}").fetchall()]
    except duckdb.Error:
        return []
    finally:
        con.close()
