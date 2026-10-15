"""Dosiér hypotézy: všetko, čo systém k hypotéze zozbieral, v jednom ZIP – čitateľný súhrn (README.md), tabuľky
(CSV a prehlad.xlsx), strojovo čitateľný dosier.json, texty zdrojových dokumentov, volania modelov, dáta Brány 1
(Parquet) a manifest so SHA-256 každého súboru.

Súvisiace hypotézy: Brána 0 (organizácie, ľudia) ostáva pri rodičovi rozdelenej hypotézy, preto sa predvolene pridá
rodič a podhypotézy; riadky nesú stĺpec `hypothesis_id`.

Súkromie (CLAUDE.md – exporty osoby predvolene pseudonymizujú): osoby ako „Osoba P<id>“, bez osobných kontaktov,
zdrojov osôb a volaní modelov o osobách; v textoch sa mená zozbieraných osôb (aj v iných pádoch – podľa kmeňa
priezviska) nahradia „[osoba]“ a e-maily a telefóny maskou. Všeobecné kontakty organizácií (podateľňa) ostávajú.
S menami a kontaktmi iba na výslovnú voľbu; taký export sa zapíše do poznámok hypotézy (čl. 5 ods. 2 GDPR)."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
import re
import subprocess
import tempfile
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from .config import Config
from .datasets import connect_ro, duck_path
from .db import DB
from .privacy import mask_text, person_pseudonym
from .registers import RPO_ATTRIBUTION
from .report import STATUS_LABEL

log = logging.getLogger("radar")

DATA_MODES = ("full", "sample", "none")
DATA_LABELS = {"full": "celé tabuľky (Parquet; do 100 tis. riadkov aj CSV)", "sample": "vzorka 1 000 riadkov (CSV)",
               "none": "bez dát (iba popis tabuliek)"}
SAMPLE_ROWS = 1000
CSV_MAX_ROWS = 100_000
XLSX_CELL_MAX = 32000
LICENSES = {"crz.gov.sk": "CC BY-SA 4.0 (Centrálny register zmlúv, Úrad vlády SR)",
            "nku.gov.sk": "CC BY 4.0 (Najvyšší kontrolný úrad SR)"}
GATE_STATUS = {"not_started": "nezačatá", "running": "prebieha", "ready": "pripravená", "needs_input": "čaká na tvoje dáta",
               "passed": "prešla", "paused": "pozastavená", "stopped": "zastavená"}
MANUAL_KINDS = {"profile_url": "profil – over ručne (LinkedIn sa nesťahuje)", "hidden_contact": "kontakt otvor ručne"}


# ======================================================================= pseudonymizácia
class Scrubber:
    """Mená zozbieraných osôb → „[osoba]“ (celé meno, priezvisko aj v iných pádoch), e-maily a telefóny → maska.
    Rozlišuje veľké písmená (mená sa píšu s veľkým začiatočným písmenom alebo celé veľkými – weby úradov)."""

    SHORT_SUFFIX = r"(?:a|ovi|om|ovia|ov|e|ou|ej|ú|ovej|ovou)?"

    def __init__(self, names: list[str]):
        patterns: set[str] = set()
        for full in names:
            parts = [p for p in re.split(r"\s+", (full or "").strip()) if len(p) > 1]
            if not parts or len(parts[-1]) < 4:
                continue
            surname = parts[-1]
            for variant in {surname, surname.upper(), surname[:1].upper() + surname[1:].lower()}:
                tail = re.escape(variant[:-1]) + r"\w*" if len(variant) >= 6 else re.escape(variant) + self.SHORT_SUFFIX
                patterns.add(tail)
                if len(parts) > 1:  # „Jana Starostová“ → celé „[osoba]“, nie „Jana [osoba]“
                    first = parts[-2]
                    for f in {first, first.upper()}:
                        patterns.add(re.escape(f) + r"\s+" + tail)
        self.re = re.compile(r"(?<!\w)(?:" + "|".join(sorted(patterns, key=len, reverse=True)) + r")(?!\w)") \
            if patterns else None

    def text(self, value: str) -> str:
        if not value:
            return value
        return mask_text(self.re.sub("[osoba]", value) if self.re else value)

    def deep(self, value: Any) -> Any:
        if isinstance(value, str):
            return self.text(value)
        if isinstance(value, dict):
            return {k: self.deep(v) for k, v in value.items()}
        if isinstance(value, list):
            return [self.deep(v) for v in value]
        return value


# ======================================================================= zber
@dataclass
class Table:
    path: str                      # cesta v ZIP bez prípony, napr. „dokazy/tvrdenia“
    title: str                     # hárok v prehlad.xlsx
    rows: list[dict[str, Any]]
    keep: tuple[str, ...] = ()     # stĺpce, ktoré pseudonymizácia nemení (všeobecné kontakty organizácií)
    csv_rows: list[dict[str, Any]] | None = None  # skrátený tvar pre CSV a XLSX (experimenty); inak rows


@dataclass
class Dossier:
    hid: int
    scope: list[tuple[int, str]]   # (id hypotézy, vzťah: hlavná / rodič / podhypotéza)
    personal: bool
    tables: list[Table] = field(default_factory=list)
    texts: dict[str, str] = field(default_factory=dict)    # cesta v ZIP → text
    files: dict[str, Path] = field(default_factory=dict)   # cesta v ZIP → súbor na disku (demá)
    summary: dict[str, Any] = field(default_factory=dict)  # podklady pre README

    @property
    def hids(self) -> list[int]:
        return [h for h, _ in self.scope]

    def table(self, path: str) -> list[dict[str, Any]]:
        return next((t.rows for t in self.tables if t.path == path), [])


def _parsed(row) -> dict[str, Any]:
    """Riadok DB → dict; stĺpce *_json sa rozbalia (v CSV sa zapíšu späť ako JSON)."""
    out = dict(row)
    for k, v in out.items():
        if k.endswith("_json") and isinstance(v, str):
            try:
                out[k] = json.loads(v)
            except ValueError:
                pass
    return out


def scope_of(db: DB, hid: int, include_related: bool = True) -> list[tuple[int, str]]:
    h = db.get_hypothesis(hid)
    if h is None:
        raise ValueError(f"Hypotéza {hid} neexistuje.")
    out = [(hid, "hlavná")]
    if include_related:
        if h["parent_id"] and db.get_hypothesis(h["parent_id"]):
            out.append((int(h["parent_id"]), "rodič"))
        out += [(int(c["id"]), "podhypotéza") for c in db.children_of(hid)]
    return out


def collect(cfg: Config, db: DB, hid: int, personal: bool = False, include_related: bool = True) -> Dossier:
    """Všetky údaje k hypotéze (a súvisiacim) ako tabuľky a texty; v režime pseudonym už očistené."""
    dos = Dossier(hid=hid, scope=scope_of(db, hid, include_related), personal=personal)
    hids = dos.hids
    relation = dict(dos.scope)
    add = dos.tables.append

    # -------------------------------------------------------------- hypotézy
    hyps = []
    for h_id in hids:
        h = _parsed(db.get_hypothesis(h_id))
        theme = db.get_theme(h["theme_id"]) if h.get("theme_id") else None
        hyps.append({"vzťah": relation[h_id], **h, "téma": theme["name"] if theme else None})
    add(Table("hypoteza/hypotezy", "Hypotézy", hyps))
    add(Table("hypoteza/poznamky", "Poznámky", [_parsed(n) for h in hids for n in db.notes_for(h)]))
    add(Table("hypoteza/kritiky", "Kritiky", [_parsed(r) for h in hids for r in db.critiques_for(h)]))
    add(Table("hypoteza/odhady_trhu", "Odhady trhu", [_parsed(r) for h in hids for r in db.market_estimates_for(h)]))

    # -------------------------------------------------------------- dôkazy
    claims = [_parsed(c) for c in db.dossier_claims(hids)]
    add(Table("dokazy/tvrdenia", "Tvrdenia", claims))
    doc_ids = sorted({c["document_id"] for c in claims})
    docs = []
    for d in db.documents_by_ids(doc_ids):
        row = _parsed(d)
        text = row.pop("text") or ""
        row["text_súbor"] = f"dokazy/texty/D{d['id']}.txt"
        docs.append(row)
        dos.texts[row["text_súbor"]] = text
    add(Table("dokazy/dokumenty", "Dokumenty", docs))

    # -------------------------------------------------------------- Brána 0
    orgs = [_parsed(o) for h in hids for o in db.orgs_for_hypothesis(h, include_excluded=True)]
    org_ids = sorted({o["id"] for o in orgs})
    people = {p["id"]: p for p in db.people_for_hypotheses(hids)}
    links = [l for h in hids for l in db.links_for_hypothesis(h)]
    contacts = {pid: db.contacts_for_person(pid) for pid in people}
    outreach = [o for h in hids for o in db.outreach_list(h)]
    interviews = [i for h in hids for i in db.interviews_for(h)]

    def who(pid: int | None) -> str:
        if pid is None or pid not in people:
            return ""
        return people[pid]["full_name"] if personal else person_pseudonym(pid)

    add(Table("brana0/organizacie", "Organizácie", orgs))
    rows = []
    linked = set()
    for l in links:
        linked.add(l["person_id"])
        rows.append(_person_row(people.get(l["person_id"]), l, contacts.get(l["person_id"], []), personal))
    for pid, p in people.items():
        if pid not in linked:
            rows.append(_person_row(p, None, contacts.get(pid, []), personal))
    add(Table("brana0/ludia", "Ľudia", rows))
    org_names = {o["id"]: o["name"] for o in orgs}
    contact_rows = [_contact_row(c, "", org_names[oid]) for oid in org_ids for c in db.org_contacts(oid)]
    if personal:
        contact_rows += [_contact_row(c, who(pid), people[pid]["org_name"]) for pid, cs in contacts.items() for c in cs]
    add(Table("brana0/kontakty", "Kontakty", contact_rows, keep=() if personal else ("hodnota", "zdroj_url")))
    add(Table("brana0/oslovenia", "Oslovenia", [
        {"id": o["id"], "hypothesis_id": o["hypothesis_id"], "osoba": who(o["person_id"]), "funkcia": o["role_title"],
         "organizácia": o["org_name"], "kanál": o["channel"], "variant": o["variant"], "stav": o["status"],
         "predmet": o["subject"], "text": o["body"], "kontakt": o["contact_value"] if personal else
         (o["contact_kind"] or ""), "schválené": o["approved_at"], "odoslané": o["sent_at"],
         "follow_up": o["follow_up_due"], "follow_up_počet": o["follow_up_count"], "fakty": _loads(o["facts_json"]),
         "kontroly": _loads(o["checks_json"]), "vytvorené": o["created_at"]} for o in outreach]))
    add(Table("brana0/rozhovory", "Rozhovory", [
        ({"meno": who(i["person_id"])} if personal else {}) | {k: i[k] for k in i.keys() if k != "person_id"}
        for i in interviews]))
    add(Table("brana0/rozhodnutia", "Rozhodnutia B0", [_parsed(r) for h in hids for r in db.decisions_for(h)]))
    role_maps = {f"H{h}": db.get_role_map(h) for h in hids if db.get_role_map(h)}

    # -------------------------------------------------------------- Brána 1
    add(Table("brana1/polia", "Polia B1", [_parsed(r) for h in hids for r in db.fields_for(h)]))
    add(Table("brana1/zdroje", "Zdroje B1", [_parsed(r) for h in hids for r in db.sources_for(h)]))
    add(Table("brana1/datasety", "Datasety", [_parsed(r) for h in hids for r in db.datasets_for(h)]))
    add(Table("brana1/pokrytie", "Pokrytie", [_parsed(r) for h in hids for r in db.coverage_for(h)]))
    add(Table("brana1/reprodukcie", "Reprodukcie", [_parsed(r) for h in hids for r in db.reproductions_for(h)]))
    add(Table("brana1/ziadosti_infozakon", "Žiadosti infozákon", [_parsed(r) for h in hids for r in db.info_requests_for(h)]))
    add(Table("brana1/vstupy_cloveka", "Vstupy človeka B1", [_parsed(r) for h in hids for r in db.gate1_inputs_for(h)]))
    add(Table("brana1/rozhodnutia", "Rozhodnutia B1", [_parsed(r) for h in hids for r in db.gate_decisions_for(h, 1)]))

    # -------------------------------------------------------------- Brána 2
    add(Table("brana2/karty", "Karty", [_parsed(r) for h in hids for r in db.cards_for(h)]))
    exps = [e for h in hids for e in db.experiments_for(h)]
    full = [_parsed(e) for e in exps]
    add(Table("brana2/experimenty", "Experimenty", full, csv_rows=[_experiment_row(e) for e in full]))
    add(Table("brana2/rozhodnutia", "Rozhodnutia B2", [_parsed(r) for h in hids for r in db.gate_decisions_for(h, 2)]))
    for e in exps:
        if e["result_json"]:
            dos.texts[f"brana2/postup_E{e['id']}.md"] = _trace(cfg, db, e)
        if e["demo_path"] and Path(e["demo_path"]).is_file():
            dos.files[f"brana2/dema/E{e['id']}_{Path(e['demo_path']).name}"] = Path(e["demo_path"])

    # -------------------------------------------------------------- proveniencia (volania modelov a označenia človeka)
    objects = {"claim": [c["id"] for c in claims], "document": doc_ids, "hypothesis": hids, "organization": org_ids,
               "experiment": [e["id"] for e in exps],
               "theme": sorted({h["theme_id"] for h in hyps if h.get("theme_id")})}
    if personal:
        objects |= {"person": sorted(people), "outreach": [o["id"] for o in outreach]}
    add(Table("proveniencia/posudky_modelov", "Posudky modelov",
              [_parsed(j) for t, ids in objects.items() for j in db.judgments_for_objects(t, ids)]))
    add(Table("proveniencia/oznacenia_cloveka", "Označenia človeka",
              [_parsed(r) for t, ids in objects.items() for r in db.labels_for_objects(t, ids)]))

    if not personal:
        names = [p["full_name"] for p in db.all_people()]
        scrub = Scrubber(names)
        for t in dos.tables:
            t.rows = [{k: (v if k in t.keep else scrub.deep(v)) for k, v in r.items()} for r in t.rows]
            if t.csv_rows is not None:
                t.csv_rows = [scrub.deep(r) for r in t.csv_rows]
        dos.texts = {k: scrub.text(v) for k, v in dos.texts.items()}
        role_maps = scrub.deep(role_maps)
        dos.summary["scrub"] = scrub
    if role_maps:
        dos.texts["brana0/mapa_roli.json"] = json.dumps(role_maps, ensure_ascii=False, indent=2, default=str)
    dos.summary |= {"people": len(people), "linked": len(linked), "person_contacts": sum(len(c) for c in contacts.values()),
                    "org_contacts": sum(1 for r in contact_rows if not r["osoba"]), "documents": len(docs)}
    return dos


def _loads(value: str | None) -> Any:
    try:
        return json.loads(value) if value else None
    except ValueError:
        return value


def _person_row(p, link, contacts: list, personal: bool) -> dict[str, Any]:
    if p is None:
        return {}
    row = {"hypothesis_id": link["hypothesis_id"] if link else p["collected_for_hypothesis_id"],
           "osoba_id": p["id"], "osoba": p["full_name"] if personal else person_pseudonym(p["id"])}
    if personal:
        row |= {"tituly": p["titles"], "rod": p["gender_hint"], "zdroj_url": p["source_url"],
                "zdroj_úryvok": p["source_snippet"], "poznámka": p["note"]}
    row |= {"funkcia": p["role_title"], "kategória_roly": p["role_category"], "organizácia": p["org_name"],
            "ičo_organizácie": p["org_ico"], "typ_organizácie": p["org_sector_kind"], "expert": p["is_expert"],
            "stav": p["status"], "druh_zdroja": p["source_kind"], "právny_základ": p["lawful_basis"],
            "väzba": link["link_type"] if link else "ešte nezoradený",
            "priorita": link["priority"] if link else None,
            "zložky_priority": _loads(link["priority_json"]) if link else None,
            "verdikt": link["verdict"] if link else None, "dôvod": link["rationale"] if link else None,
            "dôkaz_tvrdenie": link["evidence_claim_id"] if link else None,
            "dôkaz_url": link["evidence_url"] if link else None,
            "osobné_kontakty": ", ".join(f"{k} {n}" for k, n in sorted(Counter(c["kind"] for c in contacts).items())),
            "zozbieraná": p["created_at"], "vymazať_po": p["purge_after"]}
    return row


def _contact_row(c, person: str, org_name: str | None) -> dict[str, Any]:
    return {"osoba": person, "organizácia": org_name, "druh": c["kind"], "hodnota": c["value"],
            "všeobecný": bool(c["is_generic"]) or c["person_id"] is None, "poznámka": MANUAL_KINDS.get(c["kind"], ""),
            "zdroj_url": c["source_url"], "zdroj_úryvok": c["source_snippet"], "extrakcia": c["extraction"],
            "kontroly": _loads(c["checks_json"]), "pridané": c["created_at"]}


def _experiment_row(e: dict[str, Any]) -> dict[str, Any]:
    spec = e.get("spec_json") or {}
    res = e.get("result_json") or {}
    prim = res.get("hlavné") or {} if isinstance(res, dict) else {}
    concl = e.get("conclusion_json") or {}
    flags = e.get("flags_json") or {}
    return {"id": e["id"], "hypothesis_id": e["hypothesis_id"], "úloha": e.get("task") or "predpoveď (pôvodná)",
            "názov": spec.get("title", "") if isinstance(spec, dict) else "", "výrok": e["statement"], "stav": e["status"],
            "uzamknuté": e["locked_at"], "hash_predregistrácie": e["prereg_hash"], "metrika": e["metric"],
            "naivný_odhad": e["baseline"], "pravidlo_úspechu": e["success_rule"],
            "model": prim.get("model"), "hodnota_modelu": prim.get("model_hodnota"),
            "hodnota_naivného": prim.get("naivný_hodnota"), "zlepšenie_pct": prim.get("zlepšenie_pct"),
            "splnené": prim.get("splnené"),
            "verdikt": flags.get("verdikt") if isinstance(flags, dict) else None,
            "záver": ((concl.get("text") or {}).get("result_for_decision") if isinstance(concl, dict) else None),
            "zdroj_záveru": concl.get("zdroj") if isinstance(concl, dict) else None,
            "upozornenia": " | ".join(flags.get("upozornenia") or []) if isinstance(flags, dict) else "",
            "organizácia": e.get("org_name"), "vytvorené": e["created_at"]}


def _trace(cfg: Config, db: DB, exp) -> str:
    from .gate2_trace import as_markdown, explain

    try:
        return as_markdown(exp, explain(cfg, db, exp))
    except Exception as exc:  # noqa: BLE001 – postup je iba popis; chyba nesmie zastaviť export
        return f"# Experiment E{exp['id']}\n\nPostup sa nedá zostaviť: {exc}\n"


# ======================================================================= dáta Brány 1
def data_summary(cfg: Config, db: DB, hids: list[int]) -> list[dict[str, Any]]:
    """Tabuľky datasetov hypotéz (bez duplikátov – podhypotézy zdieľajú DuckDB rodiča) s počtom riadkov a stĺpcami."""
    out, seen, cons = [], set(), {}
    try:
        for h in hids:
            path = duck_path(cfg, h)
            if not path.exists():
                continue
            for d in db.datasets_for(h, only_ok=True):
                if str(path) not in cons:
                    cons[str(path)] = connect_ro(path)
                con = cons[str(path)]
                existing = {r[0] for r in con.execute("SELECT table_name FROM information_schema.tables").fetchall()}
                for t in (d["table_name"], f"{d['table_name']}_long"):
                    if t not in existing or (str(path), t) in seen:
                        continue
                    seen.add((str(path), t))
                    profile = _loads(d["profile_json"]) or {}
                    out.append({"hypothesis_id": h, "dataset_id": d["id"], "tabuľka": t, "vlastník": path.parent.name,
                                "súbor": f"{path.parent.name}_{t}", "databáza": path,
                                "riadky": con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0],
                                "stĺpce": [(c[0], c[1]) for c in con.execute(f'DESCRIBE "{t}"').fetchall()],
                                "url": d["url"], "licencia": _license(db, d), "poznámka": d["note"],
                                "okno": profile.get("časové_okno") if isinstance(profile, dict) else None,
                                "db_mb": round(path.stat().st_size / 1e6, 1)})
    finally:
        for con in cons.values():
            con.close()
    return out


def _license(db: DB, d) -> str:
    if d["source_id"]:
        src = db.get_source(d["source_id"])
        if src is not None and src["license"]:
            return src["license"]
    return next((lic for dom, lic in LICENSES.items() if dom in (d["url"] or "")), "neuvedená – over pri zdroji")


def _export_table(entry: dict[str, Any], mode: str, tmp: Path) -> list[tuple[str, Path, bool]]:
    """(cesta v ZIP, súbor, uložiť bez kompresie) pre jednu tabuľku; COPY z DuckDB iba na čítanie."""
    files = []
    con = connect_ro(entry["databáza"])
    try:
        t = entry["tabuľka"]
        base = f"brana1/data/{entry['súbor']}"
        if mode == "full":
            pq = tmp / f"{entry['súbor']}.parquet"
            con.execute(f"""COPY (SELECT * FROM "{t}") TO '{pq.as_posix()}' (FORMAT parquet, COMPRESSION zstd)""")
            files.append((base + ".parquet", pq, True))
        if mode == "sample" or (mode == "full" and entry["riadky"] <= CSV_MAX_ROWS):
            limit = f" LIMIT {SAMPLE_ROWS}" if mode == "sample" else ""
            out = tmp / f"{entry['súbor']}.csv"
            con.execute(f"""COPY (SELECT * FROM "{t}"{limit}) TO '{out.as_posix()}' (HEADER, DELIMITER ';')""")
            files.append((base + ".csv", out, False))
    finally:
        con.close()
    return files


def _schema_md(entries: list[dict[str, Any]], mode: str) -> str:
    out = ["# Dáta Brány 1", "", f"Režim: {DATA_LABELS[mode]}. Parquet otvoríš v DuckDB "
           "(`SELECT * FROM 'subor.parquet'`), Pythone (pandas.read_parquet) alebo Power BI; CSV má oddeľovač „;“ "
           "a kódovanie UTF-8.", ""]
    if not entries:
        out.append("K hypotéze nie sú stiahnuté tabuľky.")
    for e in entries:
        out += [f"## {e['súbor']}", "", f"- dataset {e['dataset_id']} (hypotéza H{e['hypothesis_id']}), "
                f"{e['riadky']:,} riadkov".replace(",", " "), f"- zdroj: {e['url']}", f"- licencia: {e['licencia']}"]
        if e.get("okno"):
            out.append(f"- časové okno: {e['okno'].get('od')} – {e['okno'].get('do')}")
        if e.get("poznámka"):
            out.append(f"- poznámka: {e['poznámka']}")
        out += ["", "| stĺpec | typ |", "|---|---|"] + [f"| {c} | {t} |" for c, t in e["stĺpce"]] + [""]
    return "\n".join(out)


# ======================================================================= zápis
def _csv_bytes(rows: list[dict[str, Any]]) -> bytes:
    cols = list(dict.fromkeys(k for r in rows for k in r))
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(cols)
    for r in rows:
        w.writerow([_cell(r.get(c)) for c in cols])
    return buf.getvalue().encode("utf-8-sig")


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, default=str)
    return v


def _xlsx_bytes(tables: list[Table]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    wb = Workbook(write_only=True)
    used = set()
    for t in tables:
        rows = t.csv_rows if t.csv_rows is not None else t.rows
        if not rows:
            continue
        title = re.sub(r"[\[\]:*?/\\]", " ", t.title)[:31]
        while title in used:
            title = title[:29] + f" {len(used)}"
        used.add(title)
        ws = wb.create_sheet(title)
        cols = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(cols)
        for r in rows:
            vals = []
            for c in cols:
                v = _cell(r.get(c))
                if isinstance(v, str):
                    v = ILLEGAL_CHARACTERS_RE.sub("", v)
                    if len(v) > XLSX_CELL_MAX:
                        v = v[:XLSX_CELL_MAX] + " … (skrátené – celé v CSV a dosier.json)"
                vals.append(v)
            ws.append(vals)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _slug(text: str, n: int = 40) -> str:
    ascii_ = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_.lower()).strip("-")[:n].strip("-") or "hypoteza"


def _commit() -> str | None:
    try:
        out = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=Path(__file__).resolve().parent.parent,
                             capture_output=True, text=True, timeout=5)
        return out.stdout.strip() or None
    except (OSError, subprocess.SubprocessError):
        return None


def write(cfg: Config, db: DB, hid: int, personal: bool = False, data_mode: str = "full", include_related: bool = True,
          progress: Callable[[str], None] | None = None) -> Path:
    """Dosiér do exports/ (mimo gitu, purge po 30 dňoch); vráti cestu k ZIP."""
    if data_mode not in DATA_MODES:
        raise ValueError(f"Dáta: {', '.join(DATA_MODES)}")
    say = progress or (lambda _m: None)
    say("Zbieram údaje z databázy")
    dos = collect(cfg, db, hid, personal, include_related)
    entries = data_summary(cfg, db, dos.hids)
    out_dir = cfg.path(cfg.exports_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    title = next(r for r in dos.table("hypoteza/hypotezy") if r["id"] == hid)["title"]
    name = f"H{hid}_{_slug(title)}_{datetime.now():%Y-%m-%d_%H%M}_{'osobne' if personal else 'pseudonym'}.zip"
    final, part = out_dir / name, out_dir / (name + ".part")
    listing: list[dict[str, Any]] = []
    try:
        with zipfile.ZipFile(part, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as z:
            def put_bytes(arc: str, data: bytes, rows: int | None = None) -> None:
                z.writestr(arc, data)
                listing.append({"cesta": arc, "bajty": len(data), "sha256": hashlib.sha256(data).hexdigest(), "riadky": rows})

            def put_file(arc: str, path: Path, stored: bool, rows: int | None = None, bom: bool = False) -> None:
                if bom:  # CSV z DuckDB: Excel potrebuje BOM, inak pokazí diakritiku
                    put_bytes(arc, b"\xef\xbb\xbf" + path.read_bytes(), rows)
                    return
                z.write(path, arc, compress_type=zipfile.ZIP_STORED if stored else zipfile.ZIP_DEFLATED)
                h = hashlib.sha256()
                with path.open("rb") as fh:
                    for chunk in iter(lambda: fh.read(1 << 20), b""):
                        h.update(chunk)
                listing.append({"cesta": arc, "bajty": path.stat().st_size, "sha256": h.hexdigest(), "riadky": rows})

            say("Zapisujem tabuľky a texty")
            for t in dos.tables:
                rows = t.csv_rows if t.csv_rows is not None else t.rows
                if rows:
                    put_bytes(t.path + ".csv", _csv_bytes(rows), len(rows))
            exps = dos.table("brana2/experimenty")
            if exps:
                put_bytes("brana2/experimenty.json", json.dumps(exps, ensure_ascii=False, indent=2, default=str).encode())
            for arc, text in sorted(dos.texts.items()):
                put_bytes(arc, (text or "").encode("utf-8"))
            scrub = dos.summary.get("scrub")
            for arc, path in dos.files.items():
                data = path.read_bytes()
                if scrub is not None and path.suffix.lower() in (".html", ".md", ".txt"):
                    data = scrub.text(data.decode("utf-8", "replace")).encode("utf-8")
                put_bytes(arc, data)
            put_bytes("prehlad.xlsx", _xlsx_bytes(dos.tables))
            put_bytes("dosier.json", json.dumps({
                "hypotéza": hid, "súvisiace": dos.scope, "režim": "osobne" if personal else "pseudonym",
                "tabuľky": {t.path: t.rows for t in dos.tables}}, ensure_ascii=False, indent=1, default=str).encode())
            if data_mode != "none" and entries:
                with tempfile.TemporaryDirectory(prefix="dosier-") as td:
                    for e in entries:
                        say(f"Exportujem dáta {e['súbor']} ({e['riadky']:,} riadkov)".replace(",", " "))
                        for arc, path, stored in _export_table(e, data_mode, Path(td)):
                            rows = min(e["riadky"], SAMPLE_ROWS) if data_mode == "sample" else e["riadky"]
                            put_file(arc, path, stored, rows, bom=arc.endswith(".csv"))
                            path.unlink()
            put_bytes("brana1/data/SCHEMA.md", _schema_md(entries, data_mode).encode("utf-8"))
            put_bytes("README.md", _readme(cfg, db, dos, entries, data_mode, listing).encode("utf-8"))
            z.writestr("manifest.json", json.dumps({
                "hypotéza": hid, "hypotézy": dos.hids, "vzťahy": dict((f"H{h}", r) for h, r in dos.scope),
                "vygenerované": datetime.now().astimezone().isoformat(timespec="seconds"),
                "režim": "osobne" if personal else "pseudonym", "dáta": data_mode, "kód": _commit(),
                "súbory": listing}, ensure_ascii=False, indent=2))
        os.replace(part, final)
    finally:
        part.unlink(missing_ok=True)
    if personal:
        db.add_note(hid, f"Export dosiéru s menami a kontaktmi: {dos.summary['people']} osôb, "
                         f"{dos.summary['person_contacts']} osobných kontaktov (súbor {final.name}).", source="export")
    log.info("Dosiér H%s: %s (%s, dáta %s, %.1f MB)", hid, final.name, "osobne" if personal else "pseudonym", data_mode,
             final.stat().st_size / 1e6)
    return final


# ======================================================================= README
def _readme(cfg: Config, db: DB, dos: Dossier, entries: list[dict[str, Any]], data_mode: str,
            listing: list[dict[str, Any]]) -> str:
    from .cards import card_text
    from .strength import evidence_strength

    hyps = {r["id"]: r for r in dos.table("hypoteza/hypotezy")}
    main = hyps[dos.hid]
    s = dos.summary
    commit = _commit()
    out = [f"# Dosiér H{dos.hid}: {main['title']}", "",
           f"Vygenerované {datetime.now():%d. %m. %Y %H:%M} · "
           f"režim: {'s menami a kontaktmi' if dos.personal else 'pseudonymizovaný (bez mien a osobných kontaktov)'} · "
           f"dáta: {DATA_LABELS[data_mode]}" + (f" · kód {commit}" if commit else ""), ""]
    if dos.personal:
        out += ["> **Obsahuje osobné údaje** (mená, pracovné kontakty zo zverejnených zdrojov). Právny základ: oprávnený "
                "záujem (LIA). Neposielaj ďalej, po použití zmaž; priečinok exports/ maže purge po 30 dňoch. "
                "Export je zapísaný v poznámkach hypotézy.", ""]
    else:
        out += ["> **Pseudonymizovaný export.** Osoby sú uvedené ako „Osoba P<id>“; mená zozbieraných osôb sú v textoch "
                "nahradené „[osoba]“, e-maily a telefóny maskou; osobné kontakty a volania modelov o osobách chýbajú. "
                "Všeobecné kontakty organizácií ostali. Nezvyčajný tvar mena sa nemusí zachytiť – pred odoslaním "
                "ďalej prezri. Citáty sú preto v tomto režime upravené; doslovné znenie je v režime s menami a v zdroji.",
                ""]
    out += ["## Hypotézy v dosiéri", "", "| H | vzťah | názov | segment | stav | test | B0 | B1 | B2 |", "|---|---|---|---|---|---|---|---|---|"]
    for h_id, rel in dos.scope:
        h = hyps[h_id]
        verdict = (h.get("validation_json") or {}).get("verdikt") if isinstance(h.get("validation_json"), dict) else None
        out.append(f"| H{h_id} | {rel} | {_md(h['title'])} | {_md(h['segment'])} | {h['status']} | {verdict or '–'} | "
                   + " | ".join(GATE_STATUS.get(h[f"gate{g}_status"], h[f"gate{g}_status"] or "–") for g in (0, 1, 2)) + " |")
    out += ["", f"## H{dos.hid} – popis", "", _md(main.get("description") or "–"), "",
            f"- pôvod: {main.get('origin')} · odvetvie: {main.get('sector') or '–'} · téma: {main.get('téma') or '–'}"]
    label, n_ver, n_part, n_dom = evidence_strength(db.evidence_for(dos.hid))
    out.append(f"- sila dôkazov: {label} (overené {n_ver}, čiastočne {n_part}, rôzne domény {n_dom})")
    g1 = main.get("gate1_json") if isinstance(main.get("gate1_json"), dict) else {}
    if g1.get("skóre") is not None:
        out.append(f"- pripravenosť dát (Brána 1): {g1['skóre']} · zložky: "
                   + ", ".join(f"{k} {v}" for k, v in (g1.get("zložky") or {}).items()))
    card = db.latest_card(dos.hid)
    if card is not None:
        out += ["", "## Problémová karta (posledná verzia)", "", card_text(json.loads(card["card_json"]))]

    claims = dos.table("dokazy/tvrdenia")
    by_status = Counter(c["status"] for c in claims)
    out += ["", "## Dôkazy", "", f"{len(claims)} tvrdení z {s['documents']} dokumentov: "
            + ", ".join(f"{STATUS_LABEL.get(k, k)} {v}" for k, v in by_status.most_common())
            + ". Všetky sú v `dokazy/tvrdenia.csv`, texty "
            "zdrojov v `dokazy/texty/`.", ""]
    for c in [c for c in claims if c["status"] == "verified" and not c.get("duplicate_of")][:10]:
        out += [f"> „{_md(c['quote'])}“", f"> — [{c['domain']}]({c['url']}), stiahnuté {str(c['fetched_at'])[:10]} "
                f"(C{c['id']}, H{c['hypothesis_id']})", ""]

    outreach = Counter(o["stav"] for o in dos.table("brana0/oslovenia"))
    out += ["## Brána 0 – organizácie a ľudia", "",
            f"- organizácie: {len({o['id'] for o in dos.table('brana0/organizacie')})}",
            f"- ľudia: {s['people']} (s väzbou na problém {s['linked']}); osobné kontakty: {s['person_contacts']}"
            + ("" if dos.personal else " – v tomto režime nie sú v exporte"),
            f"- všeobecné kontakty organizácií: {s['org_contacts']}",
            f"- oslovenia: {sum(outreach.values())}" + (f" ({', '.join(f'{k} {v}' for k, v in outreach.items())})" if outreach else ""),
            f"- rozhovory: {len(dos.table('brana0/rozhovory'))}", ""]
    if main.get("parent_id") and "rodič" not in dict(dos.scope).values():
        out += ["Ľudia a organizácie rozdelenej hypotézy ostávajú pri rodičovi – exportuj so súvisiacimi hypotézami.", ""]

    fields = dos.table("brana1/polia")
    out += ["## Brána 1 – dáta", "",
            f"- polia: {len(fields)} (kľúčové {sum(1 for f in fields if f.get('is_key'))}, nájdené "
            f"{sum(1 for f in fields if f.get('status') == 'found')}); zdroje: {len(dos.table('brana1/zdroje'))}; "
            f"reprodukcie: {len(dos.table('brana1/reprodukcie'))}", ""]
    if entries:
        out += ["| tabuľka | riadky | zdroj | licencia |", "|---|---|---|---|"]
        out += [f"| {e['súbor']} | {e['riadky']:,} | {e['url']} | {e['licencia']} |".replace(",", " ") for e in entries]
        out += ["", "Popis stĺpcov: `brana1/data/SCHEMA.md`.", ""]

    exps = dos.table("brana2/experimenty")
    out += ["## Brána 2 – experimenty", ""]
    if not exps:
        out.append("Zatiaľ žiadne.")
    for e in [_experiment_row(x) for x in exps]:
        line = f"- **E{e['id']}** (H{e['hypothesis_id']}, {e['úloha']}, {e['stav']}) {_md(e['názov'] or e['výrok'] or '')}"
        if e["hodnota_modelu"] is not None:
            line += (f" – {e['metrika']}: model {e['hodnota_modelu']} vs. naivný {e['hodnota_naivného']}, zlepšenie "
                     f"{e['zlepšenie_pct']} %, pravidlo {'splnené' if e['splnené'] else 'nesplnené'}"
                     + (f"; verdikt: {e['verdikt']}" if e.get("verdikt") else ""))
        out.append(line)
        if e["záver"]:
            out.append(f"  - záver ({e['zdroj_záveru']}): {_md(e['záver'])}")
        if e["upozornenia"]:
            out.append(f"  - upozornenia: {_md(e['upozornenia'])}")
    out += ["", "Postup každého experimentu krok po kroku: `brana2/postup_E<id>.md`.", "",
            "## Súbory", "", "| súbor | riadky | veľkosť |", "|---|---|---|"]
    out += [f"| `{f['cesta']}` | {f['riadky'] if f['riadky'] is not None else ''} | {_size(f['bajty'])} |" for f in listing]
    out += ["| `README.md`, `manifest.json` | | |", "",
            "`prehlad.xlsx` má všetky tabuľky ako hárky, `dosier.json` všetko okrem dát a textov dokumentov. "
            "Stĺpce majú názvy z databázy (`status` = stav, `*_json` = štruktúrované údaje). `manifest.json` obsahuje "
            "SHA-256 každého súboru.", "", "## Licencie a zdroje", ""]
    if any(o.get("register_json") for o in dos.table("brana0/organizacie")):
        out.append(f"- {RPO_ATTRIBUTION}")
    for lic in sorted({f"{e['url']}: {e['licencia']}" for e in entries}):
        out.append(f"- {lic}")
    out.append("- Citáty patria ich vydavateľom; sú uvedené so zdrojom na účely analýzy problému.")
    text = "\n".join(out) + "\n"
    return s["scrub"].text(text) if "scrub" in s else text  # karta a texty z DB idú do README priamo


def _md(text: Any) -> str:
    return str(text or "").replace("|", "/").replace("\n", " ").strip()


def _size(n: int) -> str:
    return f"{n / 1e6:.1f} MB" if n >= 1e6 else f"{n / 1e3:.0f} kB" if n >= 1e3 else f"{n} B"
