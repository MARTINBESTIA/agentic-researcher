"""M15: dosiér hypotézy – všetko k hypotéze v jednom ZIP (dôkazy, Brány 0–2, dáta, proveniencia).
Predvolene pseudonymizovaný: v žiadnom súbore nesmie byť meno osoby ani jej osobný kontakt."""

import hashlib
import io
import json
import zipfile

import duckdb
from openpyxl import load_workbook

from radar import dossier
from radar.config import Config, EmbeddingsConfig, Gate1Config, ModelsConfig, PeopleConfig, PrivacyConfig
from radar.datasets import hyp_dir, import_table
from radar.db import DB

DOC_TEXT = ("Starostka Jana Starostová uviedla, že obec nemá kapacitu na kontrolu zmlúv. Podľa Starostovej "
            "chýba právnik. Kontakt: starostka@obec-fikt.sk, tel. +421 905 123 456.")
SECRETS = ("Starostov", "STAROSTOV", "starostka@", "905 123", "linkedin.com/in/jana")


def setup(tmp_path):
    cfg = Config(db_path=str(tmp_path / "data" / "r.sqlite"), models=ModelsConfig(), embeddings=EmbeddingsConfig(enabled=False),
                 privacy=PrivacyConfig(lia_completed=True), people=PeopleConfig(search_profiles=False),
                 gate1=Gate1Config(nkod_enabled=False), config_dir=str(tmp_path))
    cfg.exports_dir = str(tmp_path / "exports")
    db = DB(cfg.db_path)
    parent = db.create_hypothesis("Obce nestíhajú kontrolovať zmluvy", "Široká hypotéza", "malé obce")
    child = db.create_hypothesis("Kontrola dodatkov zmlúv", "Podhypotéza", "malé obce")
    db.set_parent(child, parent)
    doc, _ = db.upsert_document("https://obce.sk/clanok", "Článok o obciach", "search", "t", "q", "h1", DOC_TEXT)
    cid = db.insert_claim(doc, 0, "t", "Obec nemá kapacitu na kontrolu zmlúv.", "malé obce", "pain",
                          "obec nemá kapacitu na kontrolu zmlúv", [])
    db.set_claim_verification(cid, "verified", 100, "citát nájdený")
    db.assign_claim(parent, cid)
    doc2, _ = db.upsert_document("https://obce.sk/dodatky", "Dodatky", "search", "t", "q", "h2",
                                 "Dodatky k zmluvám obcí pribúdajú rok po podpise.")
    cid2 = db.insert_claim(doc2, 0, "t", "Dodatky pribúdajú.", "obce", "pain", "Dodatky k zmluvám obcí pribúdajú", [])
    db.set_claim_verification(cid2, "verified", 100, "ok")
    db.assign_claim(child, cid2)
    db.add_judgment("entail", "judge", "radar-qwen3.6-27b", None, "v1", {"hypothesis": child}, {"verdict": "supported"},
                    "supported", 100, 10, 2.0, False, object_type="claim", object_id=cid2)
    db.add_note(parent, "Jana Starostová odporučila kolegu.", source="human")

    org, _ = db.upsert_organization("Obec Fiktívna", "obec fiktivna", "evidence", ico="00313271",
                                    website="https://obec-fikt.sk/", website_verified=1)
    db.link_org_hypothesis(org, parent, "mentioned_with_problem", cid, "https://obce.sk/clanok", fit={"dôkaz": "áno"})
    db.add_contact_point(None, org, "email", "podatelna@obec-fikt.sk", "https://obec-fikt.sk/kontakt",
                         "Podateľňa: podatelna@obec-fikt.sk", "mailto", is_generic=True)
    pid = db.insert_person(org, "Jana Starostová", "jana starostova", "https://obec-fikt.sk/kontakt",
                           "Jana Starostová, starostka", parent, role_title="starostka", gender_hint="female")
    ct = db.add_contact_point(pid, org, "email", "starostka@obec-fikt.sk", "https://obec-fikt.sk/kontakt",
                              "Jana Starostová, starostka@obec-fikt.sk", "mailto")
    db.add_contact_point(pid, org, "phone", "+421 905 123 456", "https://obec-fikt.sk/kontakt",
                         "Jana Starostová, +421 905 123 456", "text")
    db.add_contact_point(pid, org, "profile_url", "https://www.linkedin.com/in/jana-starostova", "https://obec-fikt.sk/",
                         "LinkedIn", "link")
    link = db.upsert_link(pid, parent, "org_affected", cid, "https://obce.sk/clanok", rationale="Jana Starostová je starostka")
    db.set_link_priority(link, 0.7, {"role_fit": 0.3})
    db.insert_outreach(pid, parent, "email", "B", "Otázka pre p. Starostovú", "Vážená pani Starostová,\ntext", [],
                       {"ok": True}, "draft", ct)
    db.insert_interview(person_id=pid, pseudonym="R1", hypothesis_id=parent, interviewee_kind="customer",
                        role_label="starostka", held_at="2026-10-01",
                        notes="Jana Starostová: kontrolu robí sama večer.")

    hyp_dir(cfg, child).mkdir(parents=True, exist_ok=True)
    csv_path = tmp_path / "zmluvy.csv"
    csv_path.write_text("cid;datum;suma\n" + "\n".join(f"{i};2025-01-{i % 28 + 1:02d};{i * 10}" for i in range(1500)) + "\n",
                        encoding="utf-8")
    import_table(cfg, child, "ds1", csv_path, "csv")
    db.add_dataset(child, None, "https://data.sk/zmluvy.csv", "ds1", status="ok", rows=1500,
                   profile_json={"časové_okno": {"od": "2025-01-01", "do": "2025-01-28"}})
    card = db.add_card(child, {"decision": {"text": "Ktoré zmluvy skontrolovať", "claims": [cid2]}})
    exp = db.add_experiment(child, None, None, "Model nájde zmluvy s dodatkom", "precision@k", "random", "aspoň o 10 %",
                            {"task": "ranking", "table": "ds1", "title": "Dodatky"}, task="ranking", card_id=card)
    db.update_experiment(exp, status="done", locked_at="2026-10-13T20:00:00+00:00", prereg_hash="abc123",
                         result_json={"hlavné": {"metrika": "precision@k", "model_hodnota": 0.4}},
                         conclusion_json={"zdroj": "kód", "text": {"result_for_decision": "Model je lepší."}})
    return cfg, db, parent, child, pid


def read_zip(path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def all_text(files: dict[str, bytes]) -> str:
    out = []
    for name, data in files.items():
        if name.endswith((".csv", ".md", ".json", ".txt")):
            out.append(data.decode("utf-8-sig"))
        elif name.endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(data), read_only=True)
            out += [str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    return "\n".join(out)


def test_pseudonymized_dossier_has_everything_but_no_personal_data(tmp_path):
    cfg, db, parent, child, pid = setup(tmp_path)
    path = dossier.write(cfg, db, child)
    assert path.parent == tmp_path / "exports" and path.name.startswith(f"H{child}_") and path.name.endswith("_pseudonym.zip")
    files = read_zip(path)
    for name in ("README.md", "prehlad.xlsx", "dosier.json", "manifest.json", "dokazy/tvrdenia.csv", "dokazy/dokumenty.csv",
                 "proveniencia/posudky_modelov.csv", "brana0/organizacie.csv", "brana0/ludia.csv", "brana0/kontakty.csv",
                 "brana0/oslovenia.csv", "brana0/rozhovory.csv", "brana1/datasety.csv", "brana1/data/SCHEMA.md",
                 "brana2/karty.csv", "brana2/experimenty.csv", "brana2/experimenty.json", "hypoteza/hypotezy.csv"):
        assert name in files, name
    text = all_text(files)
    for secret in SECRETS:  # meno (aj v inom páde), osobný e-mail, telefón, LinkedIn
        assert secret not in text, secret
    assert "podatelna@obec-fikt.sk" in files["brana0/kontakty.csv"].decode("utf-8-sig")  # všeobecný kontakt ostáva
    people = files["brana0/ludia.csv"].decode("utf-8-sig")
    assert f"Osoba P{pid}" in people and str(parent) in people  # ľudia rodiča (Brána 0 ostáva pri rodičovi)
    claims = files["dokazy/tvrdenia.csv"].decode("utf-8-sig")
    assert "https://obce.sk/clanok" in claims and "Dodatky k zmluvám obcí pribúdajú" in claims and "h2" in claims
    assert any(n.startswith("dokazy/texty/D") for n in files)
    data = [n for n in files if n.startswith("brana1/data/") and n.endswith(".parquet")]
    assert len(data) == 1 and data[0] in files and f"brana1/data/H{child}_ds1.csv" in files
    pq = tmp_path / "x.parquet"
    pq.write_bytes(files[data[0]])
    assert duckdb.sql(f"SELECT COUNT(*) FROM '{pq.as_posix()}'").fetchone()[0] == 1500
    manifest = json.loads(files["manifest.json"])
    assert manifest["režim"] == "pseudonym" and set(manifest["hypotézy"]) == {parent, child}
    for f in manifest["súbory"]:
        assert hashlib.sha256(files[f["cesta"]]).hexdigest() == f["sha256"]
    exps = json.loads(files["brana2/experimenty.json"])
    assert exps[0]["prereg_hash"] == "abc123" and exps[0]["conclusion_json"]["text"]["result_for_decision"] == "Model je lepší."
    readme = files["README.md"].decode("utf-8")
    assert "Kontrola dodatkov zmlúv" in readme and "pseudonym" in readme.lower()
    assert not any("s menami a kontaktmi" in n["text"] for n in db.notes_for(child))


def test_personal_dossier_includes_contacts_and_is_recorded(tmp_path):
    cfg, db, parent, child, pid = setup(tmp_path)
    path = dossier.write(cfg, db, child, personal=True)
    assert path.name.endswith("_osobne.zip")
    files = read_zip(path)
    contacts = files["brana0/kontakty.csv"].decode("utf-8-sig")
    assert "starostka@obec-fikt.sk" in contacts and "linkedin.com/in/jana-starostova" in contacts and "905 123 456" in contacts
    assert "Jana Starostová" in files["brana0/ludia.csv"].decode("utf-8-sig")
    assert "Starostová" in files[f"dokazy/texty/D{1}.txt"].decode("utf-8")  # pôvodný text dokumentu
    assert any("s menami a kontaktmi" in n["text"] for n in db.notes_for(child))


def test_data_modes_and_related_switch(tmp_path):
    cfg, db, parent, child, pid = setup(tmp_path)
    files = read_zip(dossier.write(cfg, db, child, data_mode="sample"))
    sample = files[f"brana1/data/H{child}_ds1.csv"].decode("utf-8-sig").strip().splitlines()
    assert len(sample) == 1 + dossier.SAMPLE_ROWS and not any(n.endswith(".parquet") for n in files)
    files = read_zip(dossier.write(cfg, db, child, data_mode="none", include_related=False))
    assert not any(n.startswith("brana1/data/") and not n.endswith("SCHEMA.md") for n in files)
    assert "brana0/ludia.csv" not in files  # bez rodiča nemá podhypotéza ľudí
    summary = dossier.data_summary(cfg, db, [child])
    assert summary[0]["riadky"] == 1500 and summary[0]["tabuľka"] == "ds1"
